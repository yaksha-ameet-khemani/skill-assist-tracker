import openpyxl, json, re, warnings, sys
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
FILE, SHEET, HR = 'Testing-Track-Master-Plan 2.xlsx', 'Track Master Plan', 4
ws = openpyxl.load_workbook(base + 'MPhasis/' + FILE, data_only=True)[SHEET]
one = lambda v: re.sub(r'\s+', ' ', str(v)).strip()
cols = [(c.column_letter, one(c.value)) for c in ws[HR] if c.value is not None]
toc, phase = [], None
for r in range(HR + 1, 53):
    a, b = ws[f'A{r}'].value, ws[f'B{r}'].value
    if a and not b: phase = one(a).split(' -- ')[0]; continue
    if not (a and b): continue
    data = {h: str(ws[f'{l}{r}'].value).strip() for l, h in cols if ws[f'{l}{r}'].value not in (None, '', '—')}
    data['Phase'] = phase
    toc.append(dict(row_number=r, topic=one(b), data=data))
by = {t['row_number']: t for t in toc}

# Tracker sheet with merged cells filled down.
cs = openpyxl.load_workbook(base + 'Details_SkillAssist.xlsx', data_only=True)['Mphasis']
g = {(c.row, c.column): c.value for row in cs.iter_rows() for c in row if c.value is not None}
for m in cs.merged_cells.ranges:
    for r in range(m.min_row + 1, m.max_row + 1): g[(r, m.min_col)] = g.get((m.min_row, m.min_col))

M = {  # item name (after "<Skill> - ") -> TOC rows
 'Repo Bootstrap': [12], 'Remotes & Branching': [12, 13], 'Production Merge Workflows': [13], 'Conflict Resolution': [13],
 'Curating the Staging Area': [12], 'Rebase for a Linear History': [13], 'Production Hotfix & Rebase Pipeline': [12, 13, 14],
 'Customer-Loan 3NF Schema Fix': [16], 'Loan-Origination ER Schema Design': [17], 'Data Hardening & Many-to-Many Modeling': [17],
 'Loan-Document Progressive Normalization (2NF → 3NF)': [16], 'BCNF Decomposition (Beyond 3NF)': [16], 'Lending Infrastructure Core': [16, 17, 18],
 'Semantic HTML Skeleton': [30], 'Flexbox Form Alignment': [31], 'Responsive 2-Column → 1-Column Grid (768px)': [31], 'Form Validation States': [31],
 'Sticky Progress-Stepper Header': [31], 'Box Model & Specificity Fix': [30], 'Interactive & Accessible Button States': [30],
 'Build "Loan Application - Step 2" (Structure-Validated)': [30, 31, 32],
 'Customers Table DDL': [20], 'Loan Applications DDL': [20], 'Repayments DDL and Index': [20], 'Seed Data DML': [20], 'Pending Over 30 Days': [21],
 'Disbursed Per Branch': [21], 'Schema Evolution with ALTER & DROP': [20], 'Outer Joins & Correlated Subqueries': [21], 'Composite Index Strategy': [21],
 'Database Comprehensive Knowledge Check': [20, 21, 22],
 'Core Class Design': [24], 'Polymorphic Inheritance': [24], 'Object Contracts': [24], 'Polymorphic Drivers': [24], 'Collection Iteration & Filtering': [25],
 'Custom Checked Exceptions': [25], 'Fault-Tolerant Loops': [25], 'High-Speed Token Parsing': [25], 'Multi-Conditional Matrices': [26],
 'Structural Guardrails': [26], 'Non-Blocking Metric Trackers': [26], 'Summary Report Ledgers': [26], 'Risk Assessment Unit Testing': [27],
 'Exception Assertions': [27], 'Maven Build Structures': [27], 'Maven Dependency Injection': [27], 'Static State Management': [24],
 'Method Overloading Variations': [24], 'Secure Resource Handling': [25], 'Eligibility Pipeline Orchestration': [26], 'Test Lifecycle Fixtures': [27],
 'Data Verification System': [24, 25, 26, 27, 28],
 'PAN Format Validation': [34], 'Destructured Contact Parsing': [34], 'Live EMI Engine DOM Manipulation': [35], 'Promise-Based Async CIBIL Check': [35],
 'DOM State UI Handlers': [35], 'Dynamic Profile Merging': [34], 'Async Form Workflow Orchestrator': [35], 'Client-Side Underwriting System': [34, 35, 36],
 'Robust Locator Strategies': [42], 'Asynchronous Explicit Waits': [42], 'Automated Form Interaction': [42], 'Navigation State Assertions': [42],
 'Page Object Model Encapsulation': [43], 'Page Object Model Encapsulation 2': [43], 'Page Object Model Encapsulation 3': [43],
 'Linear POM Chain Orchestration': [43], 'Comprehensive Flow Verification': [44], 'Negative Path Failure Testing': [44],
 'Actions Class Complex Gesture Chaining': [42], 'Automation Framework': [42, 43, 44, 45],
 'Retrofitting Type Safety': [39], 'Bounding Domains': [38, 39], 'Code Migration': [39], 'Type-Safe Financial Calculation Conversion': [39],
 'Reusable Contracts-Generic Response Wrappers': [38], 'Safe Branching-Discriminated Unions & Type Narrowing': [38], 'Type-Safe Underwriting Automation': [38, 39, 40],
 'Gherkin Happy Path Scenario Design': [47], 'Gherkin Negative Path Validation': [48], 'Data-Driven Gherkin Scenario Outlines': [47],
 'Regular Expression Step Definitions': [48], 'Step Definition POM Injection': [48], 'Gherkin Tag Organization & Suite Segmentation': [47],
 'BDD Cucumber Automation Engine': [47, 48, 49],
 'Loan Eligibility Decision Engine': [51],
}
phase_rows = {1: [8, 9, 10, 12, 13, 14, 16, 17, 18, 20, 21, 22], 2: [24, 25, 26, 27, 28], 3: [30, 31, 32, 34, 35, 36, 38, 39, 40], 4: [42, 43, 44, 45, 47, 48, 49]}
items = []
for r in range(2, cs.max_row + 1):
    name = g.get((r, 3))
    if not name: continue
    name = one(name); skill = one(g[(r, 1)]); kind = one(g[(r, 2)])
    if skill == 'Milestones':
        n = int(re.match(r'Milestone (\d+)', name).group(1))
        items.append(dict(row=r, track=skill, typ='Milestone Assessment', seq=str(n), name=name, rows=phase_rows[n]))
        continue
    key = name.split(' - ', 1)[1]
    typ = 'Capstone' if skill == 'Capstone' else kind
    items.append(dict(row=r, track=skill, typ=typ, seq=None, name=name, rows=M[key]))
used = set()
for i in items:
    used.update(i['rows'])
    print(f"{i['row']:>3} {i['track'][:12]:12} {i['typ'][:10]:10} {i['name'][:58]:58} -> {i['rows'] if len(i['rows']) < 4 else str(i['rows'][:3])[:-1] + ', …]'} {by[i['rows'][0]]['topic'][:40]}", file=sys.stderr)
print(len(toc), 'TOC rows;', len(items), 'items;', 'TOC rows without content:', [(t['row_number'], t['topic']) for t in toc if t['row_number'] not in used], file=sys.stderr)

q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
colsj = [{'letter': l, 'header': h} for l, h in cols] + [{'letter': '', 'header': 'Phase'}]
out = ['begin;', "insert into clients(name) values ('Mphasis') on conflict do nothing;"]
for tr in dict.fromkeys(i['track'] for i in items):
    out.append(f"insert into tracks(client_id,name,csm) select id,{t(tr)},'Internal' from clients where name='Mphasis' on conflict (client_id,name) do update set csm=excluded.csm;")
out.append(f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,day_column,topic_column)
select id,null,{t(FILE)},{t(SHEET)},{HR},{q(colsj)}::jsonb,null,'B' from clients where name='Mphasis'
on conflict (client_id,file_name,sheet_name) do update set columns=excluded.columns;""")
out.append(f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,null,x.topic,x.data from toc_files f join clients c on c.id=f.client_id and c.name='Mphasis', jsonb_to_recordset({q(toc)}::jsonb) as x(row_number int, topic text, data jsonb)
where f.file_name={t(FILE)} on conflict (toc_file_id,row_number) do update set topic=excluded.topic, data=excluded.data;""")
for i in items:
    src = q({'source': {'file': 'Details_SkillAssist.xlsx', 'sheet': 'Mphasis', 'row': i['row']}})
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(i['track'])} join toc_files f on f.client_id=c.id and f.file_name={t(FILE)} where c.name='Mphasis'),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'{i['typ']}',{t(i['seq'])},{t(i['name'])},'2026-07-02',{src}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date returning id)
insert into content_toc_links(content_id,toc_row_id) select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number = any(array{i['rows']}::int[]) on conflict do nothing;""")
out.append('commit;')
open('/tmp/claude-1000/-home-administrator-Downloads-IntelliJ-Projects-Skill-Assist-Tracker/a76b5546-0d98-436c-b2f1-1898ab9b6cc8/scratchpad/mph.sql', 'w').write('\n'.join(out))
