import openpyxl, json, re, warnings, sys
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
FILE, SHEET, HR = 'Program_TOC_Compressed__DSA_.xlsx', 'Day-Wise Detailed TOC', 2
ws = openpyxl.load_workbook(base + 'Capgemini/' + FILE, data_only=True)[SHEET]
keep = {'A': 'Wk', 'B': 'Day', 'C': 'Topic / Focus', 'D': 'Key Concepts Covered'}
assert [ws[f'{l}{HR}'].value for l in keep] == list(keep.values())
toc = []
for r in range(HR + 1, ws.max_row + 1):
    day, topic = ws[f'B{r}'].value, ws[f'C{r}'].value
    if not day or not topic or not re.match(r'D\d+$', str(day)): continue
    data = {h: str(ws[f'{l}{r}'].value).strip() for l, h in keep.items() if ws[f'{l}{r}'].value}
    toc.append(dict(row_number=r, day_label=str(day).strip(), topic=str(topic).strip(), data=data))
row_of_day = {int(t['day_label'][1:]): t['row_number'] for t in toc}
assert sorted(row_of_day) == list(range(1, 31)), sorted(row_of_day)

cs = openpyxl.load_workbook(base + 'Details_SkillAssist.xlsx', data_only=True)['Capgemini']
items = []
for r in (2, 3, 4):
    items.append(dict(row=r, typ='Pre-Assessment', seq=None, name=cs[f'A{r}'].value.strip(), date=None, rows=[]))
week = None
for r in range(7, cs.max_row + 1):
    typ, name = cs[f'D{r}'].value, cs[f'E{r}'].value
    if not name: continue
    a = cs[f'A{r}'].value
    m = re.match(r'W(\d+)', str(a or '')) or re.match(r'WEEK (\d+)', str(a or ''))
    if m: week = int(m.group(1))
    date = '2026-07-10' if week <= 3 else '2026-07-11'
    if typ == 'Daily Assignment':
        d = int(str(cs[f'B{r}'].value)[1:])
        items.append(dict(row=r, typ=typ, seq=str(d), name=name.strip(), date=date, rows=[row_of_day[d]]))
    else:
        n = re.search(r'(\d+)$', typ).group(1)
        days = range((week - 1) * 5 + 1, (week - 1) * 5 + 5)        # the week's 4 teaching days
        items.append(dict(row=r, typ='Milestone Assessment', seq=f'W{week} · M{n}', name=name.strip(), date=date,
                          rows=[row_of_day[d] for d in days]))
for i in items:
    print(f"{i['row']:>2} {i['typ'][:9]:9} {str(i['seq'] or ''):8} {i['name'][:60]:60} {i['date']} -> rows {i['rows']}", file=sys.stderr)
print(len(toc), 'TOC days;', len(items), 'items', file=sys.stderr)

TRACK = 'Java Full Stack + DSA (Project Orbit)'
q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
out = ['begin;', "insert into clients(name) values ('Capgemini') on conflict do nothing;",
 f"insert into tracks(client_id,name) select id,{t(TRACK)} from clients where name='Capgemini' on conflict do nothing;",
 f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,day_column,topic_column)
select c.id,tr.id,{t(FILE)},{t(SHEET)},{HR},{q([{'letter': l, 'header': h} for l, h in keep.items()])}::jsonb,'B','C'
from clients c join tracks tr on tr.client_id=c.id and tr.name={t(TRACK)} where c.name='Capgemini'
on conflict (client_id,file_name,sheet_name) do update set columns=excluded.columns, track_id=excluded.track_id;""",
 f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,x.day_label,x.topic,x.data from toc_files f, jsonb_to_recordset({q(toc)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
where f.file_name={t(FILE)}
on conflict (toc_file_id,row_number) do update set day_label=excluded.day_label, topic=excluded.topic, data=excluded.data;"""]
for i in items:
    src = q({'source': {'file': 'Details_SkillAssist.xlsx', 'sheet': 'Capgemini', 'row': i['row']}})
    dt = 'null' if i['date'] is None else f"'{i['date']}'"
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(TRACK)} join toc_files f on f.client_id=c.id and f.file_name={t(FILE)} where c.name='Capgemini'),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'{i['typ']}',{t(i['seq'])},{t(i['name'])},{dt},{src}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number = any(array{i['rows'] or [0]}::int[]) on conflict do nothing;""")
out.append('commit;')
open('/tmp/claude-1000/-home-administrator-Downloads-IntelliJ-Projects-Skill-Assist-Tracker/a76b5546-0d98-436c-b2f1-1898ab9b6cc8/scratchpad/capg.sql', 'w').write('\n'.join(out))
