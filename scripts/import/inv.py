import openpyxl, json, re, warnings, sys
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
FILE, SHEET, HR = 'Copy of Invesco_Testing_Freshers Training_V2_March 2026 2.xlsx', 'Training Program', 3
ws = openpyxl.load_workbook(base + 'Invesco/' + FILE, data_only=True)[SHEET]
one = lambda v: re.sub(r'\s+', ' ', str(v)).strip()
cols = [(c.column_letter, one(c.value)) for c in ws[HR] if c.value is not None]
toc, week_of_row, week = [], {}, None
for r in range(HR + 1, ws.max_row + 1):
    if ws[f'A{r}'].value: week = ws[f'A{r}'].value
    day, topic = ws[f'B{r}'].value, ws[f'C{r}'].value
    if not day or not topic: continue
    data = {h: str(ws[f'{l}{r}'].value).strip() for l, h in cols if ws[f'{l}{r}'].value not in (None, '')}
    data['Week'] = one(week)
    toc.append(dict(row_number=r, day_label=one(day), topic=str(topic).strip(), data=data))
    week_of_row[r] = week
def week_name(w):   # "Week 1\nSoftware Dev\n& Data Handling\n(40 Hours)" -> "Week 1: Software Dev & Data Handling"
    parts = [p.strip() for p in str(w).split('\n') if p.strip() and not re.match(r'\(\d+ Hours\)', p.strip())]
    return parts[0] + ': ' + ' '.join(parts[1:]) if re.match(r'Week \d+$', parts[0]) else ' '.join(parts)
row_of_day = {int(re.search(r'\d+', t['day_label']).group()): t['row_number'] for t in toc}
weeks = {n: week_name(week_of_row[row_of_day[n * 5]]) for n in range(1, 5)}
print(weeks, file=sys.stderr)

cs = openpyxl.load_workbook(base + 'Details_SkillAssist.xlsx', data_only=True)['Invesco']
dates = {1: '2026-07-08', 2: '2026-07-20', 3: '2026-07-31', 4: '2026-08-12'}
items = []
for r in range(2, cs.max_row + 1):
    d, topic, kind, name, loc = (cs.cell(r, c).value for c in range(1, 6))
    if not name: continue
    m = re.match(r'Milestone (\d+)', str(d))
    if m:
        n = int(m.group(1)); wk = n
        items.append(dict(row=r, typ='Milestone Assessment', seq=str(n), week=wk, name=name.strip(), loc=loc,
                          rows=[row_of_day[x] for x in range(5 * n - 4, 5 * n + 1)]))
    else:
        n = int(d); wk = (n - 1) // 5 + 1
        items.append(dict(row=r, typ='Daily Assignment', seq=str(n), week=wk, name=name.strip(), loc=loc, rows=[row_of_day[n]]))
by_row = {t['row_number']: t for t in toc}
for i in items:
    print(f"{i['row']:>2} W{i['week']} {i['typ'][:9]:9} {i['seq']:>2} {i['name'][:55]:55} -> " + ', '.join(f"{by_row[x]['day_label']}" for x in i['rows']) + f"  [{by_row[i['rows'][0]]['topic'][:40]!r}]", file=sys.stderr)
print(len(toc), 'TOC rows;', len(items), 'items', file=sys.stderr)

q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
out = ['begin;', "insert into clients(name) values ('Invesco') on conflict do nothing;"]
for n, w in weeks.items():
    out.append(f"insert into tracks(client_id,name) select id,{t(w)} from clients where name='Invesco' on conflict do nothing;")
out.append(f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,day_column,topic_column)
select id,null,{t(FILE)},{t(SHEET)},{HR},{q([{'letter': l, 'header': h} for l, h in cols])}::jsonb,'B','C' from clients where name='Invesco'
on conflict (client_id,file_name,sheet_name) do update set columns=excluded.columns;""")
out.append(f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,x.day_label,x.topic,x.data from toc_files f join clients c on c.id=f.client_id and c.name='Invesco', jsonb_to_recordset({q(toc)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
where f.file_name={t(FILE)}
on conflict (toc_file_id,row_number) do update set day_label=excluded.day_label, topic=excluded.topic, data=excluded.data;""")
for i in items:
    extra = q({'location': i['loc'], 'source': {'file': 'Details_SkillAssist.xlsx', 'sheet': 'Invesco', 'row': i['row']}})
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(weeks[i['week']])} join toc_files f on f.client_id=c.id and f.file_name={t(FILE)} where c.name='Invesco'),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'{i['typ']}',{t(i['seq'])},{t(i['name'])},'{dates[i['week']]}',{extra}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number = any(array{i['rows']}::int[]) on conflict do nothing;""")
out.append('commit;')
open('/tmp/claude-1000/-home-administrator-Downloads-IntelliJ-Projects-Skill-Assist-Tracker/a76b5546-0d98-436c-b2f1-1898ab9b6cc8/scratchpad/inv.sql', 'w').write('\n'.join(out))
