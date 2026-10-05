# TCS ION DMW sheet of Details_SkillAssist.xlsx (user-confirmed 2026-10-03): client TCS ION, track DMW Syllabus,
# CSM Priyanka, 22-06-2026. Days 1-3 are daily assignments, "Python - Milestone 1" is Milestone Assessment 1.
# The TOC was only shared as a screenshot, so it is typed in below (row numbers = the TOC sheet's rows).
import openpyxl, json, re, warnings, sys
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
FILE, SHEET, CLIENT, TRACK, CSM, DATE = 'Details_SkillAssist.xlsx', 'TCS ION DMW', 'TCS ION', 'DMW Syllabus', 'Priyanka', '2026-06-22'
TOCF = '(Added manually) DMW TOC screenshot', 'DMW Syllabus'
ws = openpyxl.load_workbook(base + FILE, data_only=True)[SHEET]
one = lambda v: re.sub(r'\s+', ' ', str(v)).strip()
# row: (week, module, sub-module, video course, proficiency)
TOC = {
 2: ('1', 'Data Preprocessing', 'Descriptive data summarization, Data cleaning, Data integration and transformation', 'Data Cleansing Master Class in Python', 'Advanced'),
 3: ('2', 'Data Warehouse', 'A multidimensional data model, Data warehouse architecture', 'Data Warehouse Concepts (DWH)', 'Intermediate'),
 4: ('3', 'Exploratory Data Analysis', 'Understanding important visualization methods to understand the hidden trends, Understanding of interactive charts, Making important inferences through visual analysis', 'Exploratory Data Analysis with Pandas and Python 3.x', 'Intermediate'),
 5: ('4', 'Association Rule Mining', 'Frequent itemset mining, Mining various kinds of association rules', 'Introduction to Data Mining', 'Intermediate'),
 6: ('5', 'Data reduction', 'Feature reduction with the help of Principal Component Analysis, Visualizing features in high dimensions (t-SNE, UMAP techniques), Factor Analysis for feature grouping, Advanced libraries for feature reduction', 'Python Machine Learning Bootcamp', 'Advanced'),
 9: ('6', 'Classification and Prediction', 'Various classification methods, Various prediction methods', 'Advanced Statistics and Data Mining for Data Science', 'Advanced'),
 10: ('7', 'Cluster Analysis', 'Partitioning methods, Hierarchical methods, Density-based methods', 'Advanced Statistics and Data Mining for Data Science', 'Advanced'),
 11: ('8', 'Outlier Analysis', 'Distance-based outlier detection, Density-based local outlier detection', None, None),
 12: ('9', 'Text Mining', 'Text retrieval methods, Text indexing techniques, Query processing techniques', 'Text Mining with Machine Learning and Python', 'Intermediate'),
 13: ('10', 'Spatial Data Mining', 'Spatial classification and spatial trend analysis, Spatial clustering methods', 'Spatial Mining-Data Mining', None),
}
H = ['Week', 'Module', 'Sub-module', 'Video Course', 'Proficiency']
toc = [dict(row_number=r, day_label=v[0], topic=v[2], data={h: x for h, x in zip(H, v) if x}) for r, v in TOC.items()]
row_of_module = {v[1].lower(): r for r, v in TOC.items()}
items, day = [], 0
for r in range(2, ws.max_row + 1):
    module, name = one(ws.cell(r, 2).value), one(ws.cell(r, 3).value)
    if module == 'Milestone Assessment':
        items.append(dict(row=r, typ='Milestone Assessment', seq='1', name=name, rows=[2, 3, 4, 5, 6]))
    else:
        day += 1
        items.append(dict(row=r, typ='Daily Assignment', seq=str(day), name=name, rows=[row_of_module[module.lower()]]))
assert ws['D2'].value == '22-06-2026'
for i in items: print(i['row'], i['typ'], i['seq'], i['name'], '->', [TOC[x][1] for x in i['rows']], file=sys.stderr)

q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
out = ['begin;', f"insert into clients(name) values ({t(CLIENT)}) on conflict do nothing;",
       f"insert into tracks(client_id,name,csm) select id,{t(TRACK)},{t(CSM)} from clients where name={t(CLIENT)} on conflict (client_id,name) do update set csm=excluded.csm;",
       f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,day_column,topic_column,notes)
select c.id,tr.id,{t(TOCF[0])},{t(TOCF[1])},1,{q([{'letter': l, 'header': h} for l, h in zip('ABCDF', H)])}::jsonb,'A','C',{t('Typed in from the screenshot shared on 2026-10-03; no Excel TOC available.')} from clients c join tracks tr on tr.client_id=c.id and tr.name={t(TRACK)} where c.name={t(CLIENT)}
on conflict (client_id,file_name,sheet_name) do update set columns=excluded.columns;""",
       f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,x.day_label,x.topic,x.data from toc_files f join clients c on c.id=f.client_id and c.name={t(CLIENT)}, jsonb_to_recordset({q(toc)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
where f.file_name={t(TOCF[0])} and f.sheet_name={t(TOCF[1])}
on conflict (toc_file_id,row_number) do update set day_label=excluded.day_label, topic=excluded.topic, data=excluded.data;"""]
for i in items:
    extra = q({'source': {'file': FILE, 'sheet': SHEET, 'row': i['row']}})
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(TRACK)} join toc_files f on f.client_id=c.id and f.file_name={t(TOCF[0])} and f.sheet_name={t(TOCF[1])} where c.name={t(CLIENT)}),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'{i['typ']}',{t(i['seq'])},{t(i['name'])},'{DATE}',{extra}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number = any(array{i['rows']}::int[]) on conflict do nothing;""")
out.append('commit;')
print('\n'.join(out))
