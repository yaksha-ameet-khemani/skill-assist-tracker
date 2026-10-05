# CTRLS sheet of Details_SkillAssist.xlsx (user-confirmed 2026-10-03): track Oracle, CSM Priyanka,
# 2 Milestone Assessments (milestone 2, merged cells), 15-06-2026, one manual topic each (no TOC file).
import openpyxl, json, re, warnings, sys
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
FILE, SHEET, CLIENT, TRACK, CSM = 'Details_SkillAssist.xlsx', 'CTRLS', 'CTRLS', 'Oracle', 'Priyanka'
MANUAL = '(Added manually)', 'CTRLS topics'
ws = openpyxl.load_workbook(base + FILE, data_only=True)[SHEET]
one = lambda v: re.sub(r'\s+', ' ', str(v)).strip()
assert one(ws['A2'].value) == TRACK and str(ws['B2'].value) == '2' and ws['D2'].value == '15-06-2026'
topics = {2: 'SQL Schema Design, Constraints & DML', 3: 'SQL Joins, Aggregation & Subqueries'}
items = [dict(row=r, name=one(ws.cell(r, 3).value), rows=[r]) for r in topics]
mtoc = [dict(row_number=r, day_label=None, topic=tp, data={'Topic': tp}) for r, tp in topics.items()]
for i in items: print(i['row'], i['name'], '->', topics[i['row']], file=sys.stderr)

q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
out = ['begin;', f"insert into clients(name) values ({t(CLIENT)}) on conflict do nothing;",
       f"insert into tracks(client_id,name,csm) select id,{t(TRACK)},{t(CSM)} from clients where name={t(CLIENT)} on conflict (client_id,name) do update set csm=excluded.csm;",
       f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns)
select id,null,{t(MANUAL[0])},{t(MANUAL[1])},1,{q([{'letter': 'A', 'header': 'Topic'}])}::jsonb from clients where name={t(CLIENT)}
on conflict (client_id,file_name,sheet_name) do nothing;""",
       f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,x.day_label,x.topic,x.data from toc_files f join clients c on c.id=f.client_id and c.name={t(CLIENT)}, jsonb_to_recordset({q(mtoc)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
where f.file_name={t(MANUAL[0])} and f.sheet_name={t(MANUAL[1])}
on conflict (toc_file_id,row_number) do update set topic=excluded.topic, data=excluded.data;"""]
for i in items:
    extra = q({'source': {'file': FILE, 'sheet': SHEET, 'row': i['row']}})
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(TRACK)} join toc_files f on f.client_id=c.id and f.file_name={t(MANUAL[0])} and f.sheet_name={t(MANUAL[1])} where c.name={t(CLIENT)}),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'Milestone Assessment','2',{t(i['name'])},'2026-06-15',{extra}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number = any(array{i['rows']}::int[]) on conflict do nothing;""")
out.append('commit;')
print('\n'.join(out))
