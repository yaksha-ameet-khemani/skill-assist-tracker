# Amdocs sheet of Details_SkillAssist.xlsx (user-approved 2026-10-03).
# Rows 2-7 (13-02-2026, CSM Muzzamil): tracks Java / SQL from the Skill column, Milestone Assessment, day blank,
#   topic = Course Name. Rows 9-23 (09-03-2026, CSM Aarti): track "Java Fundamentals", topic column = TOC;
#   topic items = Daily Assignment Day 1-13 in row order, "Milestone Assessment N" rows = Milestone Assessment
#   linked to the topics since the previous milestone.
import openpyxl, json, re, warnings, sys
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
FILE, SHEET, CLIENT = 'Details_SkillAssist.xlsx', 'Amdocs', 'Amdocs'
one = lambda v: None if v is None or not str(v).strip() else re.sub(r'\s+', ' ', str(v)).strip()
iso = lambda d: '-'.join(reversed(d.split('-')))
ws = openpyxl.load_workbook(base + FILE, data_only=True)[SHEET]
toc, items, csm = [], [], {}

skill = course_row = None; date = iso(ws['D2'].value)
for r in range(2, 8):
    if one(ws[f'A{r}'].value): skill = one(ws[f'A{r}'].value)
    if one(ws[f'B{r}'].value):
        course_row = r
        toc.append(dict(row_number=r, day_label=None, topic=one(ws[f'B{r}'].value), data={'Skill': skill, 'Course Name': one(ws[f'B{r}'].value)}))
    csm[skill] = 'Muzzamil'
    items.append(dict(row=r, track=skill, typ='Milestone Assessment', seq=None, name=one(ws[f'C{r}'].value), date=date, rows=[course_row]))

TRACK = 'Java Fundamentals'; csm[TRACK] = 'Aarti'
date, day, since_m = iso(ws['D9'].value), 0, []
for r in range(9, 24):
    b, name = one(ws[f'B{r}'].value), one(ws[f'C{r}'].value)
    m = re.match(r'Milestone Assessment (\d+)$', b)
    if m:
        items.append(dict(row=r, track=TRACK, typ='Milestone Assessment', seq=m.group(1), name=name, date=date, rows=since_m))
        since_m = []
        continue
    day += 1; since_m.append(r)
    toc.append(dict(row_number=r, day_label=str(day), topic=b, data={'Skill': 'Java', 'Topic': b}))
    items.append(dict(row=r, track=TRACK, typ='Daily Assignment', seq=str(day), name=name, date=date, rows=[r]))
by_row = {t['row_number']: t['topic'] for t in toc}
for i in items: print(f"{i['row']:>2} {i['track'][:17]:17} {csm[i['track']]:8} {i['typ'][:9]} {i['seq'] or '-':>2} {i['date']} {i['name'][:45]:45} -> {[by_row[x][:22] for x in i['rows']]}", file=sys.stderr)
print(len(toc), 'topics;', len(items), 'items', file=sys.stderr)

q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
out = ['begin;', f"insert into clients(name) values ({t(CLIENT)}) on conflict do nothing;"]
for tr in dict.fromkeys(i['track'] for i in items):
    out.append(f"insert into tracks(client_id,name,csm) select id,{t(tr)},{t(csm[tr])} from clients where name={t(CLIENT)} on conflict (client_id,name) do update set csm=excluded.csm;")
cols = [{'letter': 'A', 'header': 'Skill'}, {'letter': 'B', 'header': 'Course Name'}, {'letter': 'B', 'header': 'Topic'}]
out.append(f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,topic_column,notes)
select id,null,{t(FILE)},{t(SHEET)},1,{q(cols)}::jsonb,'B',{t('Rows 2-7: Course Name as topic. Rows 9-23: column B as topic (day numbers assigned in row order).')} from clients where name={t(CLIENT)}
on conflict (client_id,file_name,sheet_name) do update set columns=excluded.columns, notes=excluded.notes;""")
out.append(f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,x.day_label,x.topic,x.data from toc_files f join clients c on c.id=f.client_id and c.name={t(CLIENT)}, jsonb_to_recordset({q(toc)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
where f.file_name={t(FILE)} and f.sheet_name={t(SHEET)}
on conflict (toc_file_id,row_number) do update set day_label=excluded.day_label, topic=excluded.topic, data=excluded.data;""")
for i in items:
    extra = q({'source': {'file': FILE, 'sheet': SHEET, 'row': i['row']}})
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(i['track'])} join toc_files f on f.client_id=c.id and f.file_name={t(FILE)} and f.sheet_name={t(SHEET)} where c.name={t(CLIENT)}),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'{i['typ']}',{t(i['seq'])},{t(i['name'])},'{i['date']}',{extra}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number = any(array{i['rows']}::int[]) on conflict do nothing;""")
out.append('commit;')
print('\n'.join(out))
