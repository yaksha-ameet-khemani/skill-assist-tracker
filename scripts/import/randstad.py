# Randstad sheet of Details_SkillAssist.xlsx, rows 9-36 only (user-confirmed 2026-10-03), CSM Aarti.
# Rows 10-25: one program, Modules 1-4, one track per module. Topic = Topic column, subtopics = Subtopics column.
#   Assignment-N -> Daily Assignment (Day 1.. in row order per track), Assessment-N -> Milestone Assessment;
#   each item linked to its own row.
# Row 27: track "AI-Augmented Systems Engineer" (full title in track notes), Milestone Assessment, day blank.
# Rows 30-36: per-participant assessments -> track "Individual Assessments", type Assessment, Skill = topic,
#   participant name kept in extra.participant (Participant column on the home page). Row 36 continues row 35
#   (Abhinav Chaudhary, GCP – Generative AI Leader, 02-09-2026).
import openpyxl, json, re, warnings, sys
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
FILE, SHEET, CLIENT, CSM = 'Details_SkillAssist.xlsx', 'Randstad', 'Randstad', 'Aarti'
MODULE_TRACK = {'Module 1': 'Module 1: ReactJs', 'Module 2': 'Module 2: NextJs',
                'Module 3': 'Module 3: React Native', 'Module 4': 'Module 4: React Navigation & State'}
AI_TRACK, IND_TRACK = 'AI-Augmented Systems Engineer', 'Individual Assessments'
one = lambda v: None if v is None or not str(v).strip() else re.sub(r'\s+', ' ', str(v)).strip()
iso = lambda d: '-'.join(reversed(d.split('-')))
ws = openpyxl.load_workbook(base + FILE, data_only=True)[SHEET]
toc, items, notes = [], [], {}

# Modules 1-4 (rows 10-25). Rows 10-18 have no module in col A: ReactJs = Module 1, NextJs = Module 2 (user).
module = topic = date = None; day = {}
for r in range(10, 26):
    name = one(ws[f'E{r}'].value)
    if not name: continue
    if one(ws[f'A{r}'].value): module = one(ws[f'A{r}'].value)
    if one(ws[f'B{r}'].value): topic = one(ws[f'B{r}'].value)
    if r <= 18: module = 'Module 1' if r <= 13 else 'Module 2'
    if ws[f'F{r}'].value: date = iso(ws[f'F{r}'].value)
    track = MODULE_TRACK[module]
    toc.append(dict(row_number=r, day_label=None, topic=topic, data={'Module': module, 'Topic': topic, 'Subtopics': one(ws[f'C{r}'].value)}))
    if one(ws[f'D{r}'].value).startswith('Assessment'):
        items.append(dict(row=r, track=track, typ='Milestone Assessment', seq='1', name=name, date=date, rows=[r]))
    else:
        day[track] = day.get(track, 0) + 1
        items.append(dict(row=r, track=track, typ='Daily Assignment', seq=str(day[track]), name=name, date=date, rows=[r]))

# Row 27
r = 27
notes[AI_TRACK] = 'Sheet title: ' + one(ws[f'A{r}'].value)
toc.append(dict(row_number=r, day_label=None, topic=one(ws[f'B{r}'].value),
                data={'Module': one(ws[f'A{r}'].value), 'Topic': one(ws[f'B{r}'].value), 'Subtopics': one(ws[f'C{r}'].value)}))
items.append(dict(row=r, track=AI_TRACK, typ='Milestone Assessment', seq=None, name=one(ws[f'E{r}'].value), date=iso(ws[f'F{r}'].value), rows=[r]))

# Rows 30-36: column A = participant (merged cells carry down).
skill_row = person = None; date = iso(ws['D30'].value)
for r in range(30, 37):
    name, skill = one(ws[f'C{r}'].value), one(ws[f'B{r}'].value)
    if one(ws[f'A{r}'].value): person = one(ws[f'A{r}'].value)
    if skill:
        skill_row = r
        toc.append(dict(row_number=r, day_label=None, topic=skill, data={'Skill': skill}))
    items.append(dict(row=r, track=IND_TRACK, typ='Assessment', seq=None, name=name, date=date, rows=[skill_row], participant=person))

by_row = {t['row_number']: t['topic'] for t in toc}
for i in items: print(f"{i['row']:>2} {i['track'][:22]:22} {i['typ'][:10]:10} {i['seq'] or '-':>2} {i['date']} {i['name'][:50]:50} -> {[by_row[x][:30] for x in i['rows']]}", file=sys.stderr)
print(len(toc), 'topics;', len(items), 'items', file=sys.stderr)

q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
out = ['begin;', f"insert into clients(name) values ({t(CLIENT)}) on conflict do nothing;"]
for tr in dict.fromkeys(i['track'] for i in items):
    out.append(f"insert into tracks(client_id,name,csm,notes) select id,{t(tr)},{t(CSM)},{t(notes.get(tr))} from clients where name={t(CLIENT)} on conflict (client_id,name) do update set csm=excluded.csm, notes=excluded.notes;")
cols = [{'letter': 'A', 'header': 'Module'}, {'letter': 'B', 'header': 'Topic'}, {'letter': 'C', 'header': 'Subtopics'}, {'letter': 'B', 'header': 'Skill'}]
out.append(f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,topic_column,notes)
select id,null,{t(FILE)},{t(SHEET)},9,{q(cols)}::jsonb,'B',{t('Rows 10-27: Module/Topic/Subtopics (header row 9). Rows 30-36: Skill (header row 29); participant kept on each item.')} from clients where name={t(CLIENT)}
on conflict (client_id,file_name,sheet_name) do update set columns=excluded.columns, notes=excluded.notes;""")
out.append(f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,x.day_label,x.topic,x.data from toc_files f join clients c on c.id=f.client_id and c.name={t(CLIENT)}, jsonb_to_recordset({q(toc)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
where f.file_name={t(FILE)} and f.sheet_name={t(SHEET)}
on conflict (toc_file_id,row_number) do update set topic=excluded.topic, data=excluded.data;""")
for i in items:
    extra = q({**({'participant': i['participant']} if i.get('participant') else {}), 'source': {'file': FILE, 'sheet': SHEET, 'row': i['row']}})
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(i['track'])} join toc_files f on f.client_id=c.id and f.file_name={t(FILE)} and f.sheet_name={t(SHEET)} where c.name={t(CLIENT)}),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'{i['typ']}',{t(i['seq'])},{t(i['name'])},'{i['date']}',{extra}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number = any(array{i['rows']}::int[]) on conflict do nothing;""")
out.append('commit;')
print('\n'.join(out))
