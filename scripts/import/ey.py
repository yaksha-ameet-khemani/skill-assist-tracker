# NOTE: the two week tracks this created were merged into one track "AI & Data" directly in the DB (2026-10-03).
# EY GDS sheet of Details_SkillAssist.xlsx, rows 1-22 (user-confirmed 2026-10-03): CSM Muzzamil, tracks = TOC week
# headings. TOC = EY GDS/Techademy_Updated AI & Data_Feasibility.xlsx "Detailed Curriculum Final" (all 6 weeks),
# only Daily Topic (topic) + Coverage (subtopics) kept. Its rows line up 1:1 with the content sheet, so a merged
# assignment cell links to every TOC row it spans; C17 (Middleware) also covers row 18. Days numbered per week.
import openpyxl, json, re, warnings, sys
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
TFILE, TSHEET = 'Techademy_Updated AI & Data_Feasibility.xlsx', 'Detailed Curriculum Final'
FILE, SHEET, CLIENT, CSM, LAST = 'Details_SkillAssist.xlsx', 'EY GDS', 'EY GDS', 'Muzzamil', 22
one = lambda v: None if v is None or not str(v).strip() else re.sub(r'\s+', ' ', str(v)).strip()
lines = lambda v: '\n'.join(re.sub(r'^\s*-\s*', '', l).strip() for l in str(v).strip().split('\n') if l.strip())
skip = re.compile(r'(pre-assessment|milestone assessment|final assessment)$', re.I)

tw = openpyxl.load_workbook(base + 'EY GDS/' + TFILE, data_only=True)[TSHEET]
toc = []
for r in range(2, tw.max_row + 1):
    topic, cov = one(tw[f'B{r}'].value), tw[f'C{r}'].value
    if not topic or skip.match(topic): continue
    data = {'Daily Topic': topic}
    if cov: data['Coverage'] = lines(cov)
    toc.append(dict(row_number=r, day_label=None, topic=topic, data=data))
toc_topic = {t['row_number']: t['topic'] for t in toc}

ws = openpyxl.load_workbook(base + FILE, data_only=True)[SHEET]
assert all(one(ws[f'B{r}'].value) == one(tw[f'B{r}'].value) for r in range(2, LAST + 1)), 'rows no longer line up'
span = {}                                    # first row of a merged C cell -> rows it covers
for m in ws.merged_cells.ranges:
    if m.min_col == 3: span[m.min_row] = list(range(m.min_row, m.max_row + 1))
span[17] = [17, 18]
items, track, date, day, wrows = [], None, None, 0, []
for r in range(2, LAST + 1):
    if ws[f'A{r}'].value: track, day, wrows = one(ws[f'A{r}'].value), 0, []
    if ws[f'D{r}'].value: date = '-'.join(reversed(ws[f'D{r}'].value.split('-')))
    if r in toc_topic: wrows.append(r)
    name = one(ws[f'C{r}'].value)
    if not name: continue
    if one(ws[f'B{r}'].value) == 'Milestone Assessment':
        n = re.search(r'(\d+)', name).group(1)
        items.append(dict(row=r, track=track, typ='Milestone Assessment', seq=n, name=name, date=date, rows=list(wrows)))
    else:
        day += 1
        items.append(dict(row=r, track=track, typ='Daily Assignment', seq=str(day), name=name, date=date, rows=span.get(r, [r])))
for i in items: print(f"{i['row']:>2} {i['track'][:7]} {i['typ'][:9]:9} {i['seq']:>2} {i['date']} {i['name'][:55]:55} -> {[toc_topic[x][:25] for x in i['rows']]}", file=sys.stderr)
print(len(toc), 'TOC rows;', len(items), 'items', file=sys.stderr)

q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
out = ['begin;', f"insert into clients(name) values ({t(CLIENT)}) on conflict do nothing;"]
for tr in dict.fromkeys(i['track'] for i in items):
    out.append(f"insert into tracks(client_id,name,csm) select id,{t(tr)},{t(CSM)} from clients where name={t(CLIENT)} on conflict (client_id,name) do update set csm=excluded.csm;")
out.append(f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,day_column,topic_column)
select id,null,{t(TFILE)},{t(TSHEET)},1,{q([{'letter': 'B', 'header': 'Daily Topic'}, {'letter': 'C', 'header': 'Coverage'}])}::jsonb,null,'B' from clients where name={t(CLIENT)}
on conflict (client_id,file_name,sheet_name) do update set columns=excluded.columns;""")
out.append(f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,x.day_label,x.topic,x.data from toc_files f join clients c on c.id=f.client_id and c.name={t(CLIENT)}, jsonb_to_recordset({q(toc)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
where f.file_name={t(TFILE)} and f.sheet_name={t(TSHEET)}
on conflict (toc_file_id,row_number) do update set day_label=excluded.day_label, topic=excluded.topic, data=excluded.data;""")
for i in items:
    extra = q({'source': {'file': FILE, 'sheet': SHEET, 'row': i['row']}})
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(i['track'])} join toc_files f on f.client_id=c.id and f.file_name={t(TFILE)} and f.sheet_name={t(TSHEET)} where c.name={t(CLIENT)}),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'{i['typ']}',{t(i['seq'])},{t(i['name'])},'{i['date']}',{extra}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number = any(array{i['rows']}::int[]) on conflict do nothing;""")
out.append('commit;')
print('\n'.join(out))
