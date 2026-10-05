# Netcracker sheet of Details_SkillAssist.xlsx, skill blocks with levels L2-L5 x 3 assignments (user-confirmed 2026-10-04).
# Each block: track = skill (column B; "Demo - <skill>" for demo blocks, which also get extra.demo = true), all Daily
# Assignment, date from column F (merged per block). Column C (L2-L5) -> Proficiency column (extra.proficiency), column B
# -> Course column (extra.course). No TOC: one manual topic per item, Module = skill.
#   rows 11-22: Demo - Java, CSM Aditya          rows 24-61: PostgreSQL, PLSQL, C++, CSM Aarti (not demo)
import openpyxl, json, re, warnings, sys
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
FILE, SHEET, CLIENT = 'Details_SkillAssist.xlsx', 'Netcracker', 'Netcracker'
BLOCKS = [(range(11, 23), 'Aditya', True), (range(24, 36), 'Aarti', False), (range(37, 49), 'Aarti', False), (range(50, 62), 'Aarti', False)]
MANUAL = '(Added manually)', 'Netcracker topics'
TOPICS = {11: 'Strings & inheritance', 12: 'Operators & control flow', 13: 'OOP – abstraction & polymorphism',
          14: 'Operators & conditionals', 15: 'Recursion', 16: 'String validation',
          17: 'Parallel streams & memory optimization', 18: 'Reflection & custom annotations',
          19: 'Concurrency – executors & thread pools', 20: 'Atomic variables (lock-free)',
          21: 'AtomicReference & CAS', 22: 'ABA problem & AtomicStampedReference',
          24: 'Joins & GROUP BY aggregation', 25: 'Transactions (COMMIT / ROLLBACK)', 26: 'Table design, keys & constraints',
          27: 'CTEs & window functions', 28: 'Triggers & PL/pgSQL', 29: 'Table partitioning & index optimization',
          30: 'Slow-query monitoring (pg_stat_statements)', 31: 'Advanced indexing for search', 32: 'Row-level security & access control',
          33: 'Foreign data wrappers & sharding', 34: 'MVCC & transaction visibility', 35: 'Buffer cache & WAL monitoring',
          37: 'PL/SQL blocks & variables', 38: 'Explicit cursors', 39: 'Stored procedures',
          40: 'Records & collections', 41: 'Packages (procedures & functions)', 42: 'Triggers & dynamic SQL',
          43: 'Bind variables & performance', 44: 'Deterministic functions & result cache', 45: 'Pipelined functions & exception handling',
          46: 'Object dependencies', 47: 'Code instrumentation & debug logging', 48: 'Call stack debugging',
          50: 'Strings & concatenation', 51: 'Basic logic & conditionals', 52: 'String traversal & character checks',
          53: 'Classes & STL containers', 54: 'String validation', 55: 'Loops & arithmetic',
          56: 'Even/odd logic & transformations', 57: 'Lambdas', 58: 'Templates & generic lambdas',
          59: 'std::atomic & multithreading', 60: 'Pointers & undefined behavior', 61: 'Loop optimization & vectorization (SIMD)'}
ws = openpyxl.load_workbook(base + FILE, data_only=True)[SHEET]
one = lambda v: re.sub(r'\s+', ' ', str(v)).strip()
assert one(ws['A11'].value) == 'Demo'
items = []
for rows, csm, demo in BLOCKS:
    course = one(ws[f'B{rows[0]}'].value); d, m, y = one(ws[f'F{rows[0]}'].value).split('-'); prof = None
    for r in rows:
        prof = one(ws[f'C{r}'].value) if ws[f'C{r}'].value else prof
        items.append(dict(row=r, name=one(ws[f'E{r}'].value), course=course, prof=prof, csm=csm, demo=demo,
                          track=('Demo - ' if demo else '') + course, date=f'{y}-{m}-{d}'))
toc = [dict(row_number=i['row'], day_label=None, topic=TOPICS[i['row']],
            data={'Module': i['course'], 'Proficiency': i['prof'], 'Topic': TOPICS[i['row']]}) for i in items]
for i in items: print(i['row'], i['track'], i['csm'], i['date'], i['prof'], i['name'], '->', TOPICS[i['row']], file=sys.stderr)

q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
cols = [{'letter': 'A', 'header': 'Module'}, {'letter': 'B', 'header': 'Proficiency'}, {'letter': 'C', 'header': 'Topic'}]
out = ['begin;']
for tr, csm in dict.fromkeys((i['track'], i['csm']) for i in items):
    out.append(f"insert into tracks(client_id,name,csm) select id,{t(tr)},{t(csm)} from clients where name={t(CLIENT)} on conflict (client_id,name) do update set csm=excluded.csm;")
out += [f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,topic_column,notes)
select c.id,null,{t(MANUAL[0])},{t(MANUAL[1])},1,{q(cols)}::jsonb,'C','No TOC for the L2-L5 skill blocks; one topic per item written by hand.'
from clients c where c.name={t(CLIENT)} on conflict (client_id,file_name,sheet_name) do update set track_id=null, notes=excluded.notes;""",
       f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,x.day_label,x.topic,x.data from toc_files f join clients c on c.id=f.client_id and c.name={t(CLIENT)}, jsonb_to_recordset({q(toc)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
where f.file_name={t(MANUAL[0])} and f.sheet_name={t(MANUAL[1])}
on conflict (toc_file_id,row_number) do update set topic=excluded.topic, data=excluded.data;"""]
for i in items:
    extra = {'source': {'file': FILE, 'sheet': SHEET, 'row': i['row']}, 'course': i['course'], 'proficiency': i['prof']}
    if i['demo']: extra['demo'] = True
    extra = q(extra)
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(i['track'])} join toc_files f on f.client_id=c.id and f.file_name={t(MANUAL[0])} and f.sheet_name={t(MANUAL[1])} where c.name={t(CLIENT)}),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'Daily Assignment',null,{t(i['name'])},'{i['date']}',{extra}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number={i['row']} on conflict do nothing;""")
out.append('commit;')
print('\n'.join(out))
