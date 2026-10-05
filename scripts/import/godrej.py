# Godrej sheet of Details_SkillAssist.xlsx (user-confirmed 2026-10-03): CSM Priyanka, all Milestone Assessments
# (day blank), 13-05-2026. Tracks renamed by the user: "React with .NET" and "Django with PostgreSQL" (the sheet's
# full titles are kept in tracks.notes). No TOC: one manual topic per item, the sheet's skill column as Module.
import openpyxl, json, re, warnings, sys
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
FILE, SHEET, CLIENT, CSM, DATE = 'Details_SkillAssist.xlsx', 'Godrej', 'Godrej', 'Priyanka', '2026-05-13'
MANUAL = '(Added manually)', 'Godrej topics'
TRACKS = {'React Bootcamp with .NET API (Router, Redux Toolkit, Hooks)': 'React with .NET',
          'Building Web Applications with Django and PostgreSQL': 'Django with PostgreSQL'}
TOPICS = {2: 'React State', 3: 'Hooks in React', 4: '.NET API CRUD', 5: '.NET API Authentication & Authorization',
          6: 'Django Models, ORM & PostgreSQL', 7: 'Django Forms, Views & Templates'}
ws = openpyxl.load_workbook(base + FILE, data_only=True)[SHEET]
one = lambda v: re.sub(r'\s+', ' ', str(v)).strip()
assert ws['D2'].value == '13-05-2026'
items, toc, full = [], [], None
for r in TOPICS:
    if ws[f'A{r}'].value: full = one(ws[f'A{r}'].value)
    skill = one(ws[f'B{r}'].value)
    toc.append(dict(row_number=r, day_label=None, topic=TOPICS[r], data={'Module': skill, 'Topic': TOPICS[r]}))
    items.append(dict(row=r, track=TRACKS[full], full=full, name=one(ws[f'C{r}'].value), rows=[r]))
for i in items: print(i['row'], i['track'], '|', i['name'], '->', TOPICS[i['row']], file=sys.stderr)

q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
out = ['begin;', f"insert into clients(name) values ({t(CLIENT)}) on conflict do nothing;"]
for full, tr in TRACKS.items():
    out.append(f"insert into tracks(client_id,name,csm,notes) select id,{t(tr)},{t(CSM)},{t('Sheet title: ' + full)} from clients where name={t(CLIENT)} on conflict (client_id,name) do update set csm=excluded.csm, notes=excluded.notes;")
out.append(f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns)
select id,null,{t(MANUAL[0])},{t(MANUAL[1])},1,{q([{'letter': 'A', 'header': 'Module'}, {'letter': 'B', 'header': 'Topic'}])}::jsonb from clients where name={t(CLIENT)}
on conflict (client_id,file_name,sheet_name) do update set columns=excluded.columns;""")
out.append(f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,x.day_label,x.topic,x.data from toc_files f join clients c on c.id=f.client_id and c.name={t(CLIENT)}, jsonb_to_recordset({q(toc)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
where f.file_name={t(MANUAL[0])} and f.sheet_name={t(MANUAL[1])}
on conflict (toc_file_id,row_number) do update set topic=excluded.topic, data=excluded.data;""")
for i in items:
    extra = q({'source': {'file': FILE, 'sheet': SHEET, 'row': i['row']}})
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(i['track'])} join toc_files f on f.client_id=c.id and f.file_name={t(MANUAL[0])} and f.sheet_name={t(MANUAL[1])} where c.name={t(CLIENT)}),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'Milestone Assessment',null,{t(i['name'])},'{DATE}',{extra}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number = any(array{i['rows']}::int[]) on conflict do nothing;""")
out.append('commit;')
print('\n'.join(out))
