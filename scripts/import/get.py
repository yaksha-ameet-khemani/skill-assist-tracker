import openpyxl, json, warnings
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
FILE, SHEET, HR = 'Sony_GET_FY2026_Plan_v5.xlsx', 'Track A — Java FSD', 5
ws = openpyxl.load_workbook(base + 'Sony GET 2026/' + FILE, data_only=True)[SHEET]
cols = [(c.column_letter, str(c.value).strip()) for c in ws[HR] if c.value is not None]
toc = []
for r in range(HR + 1, ws.max_row + 1):
    week, topic = ws[f'A{r}'].value, ws[f'D{r}'].value
    if not week or not topic or not ws[f'C{r}'].value: continue   # module rows only
    data = {h: str(ws[f'{l}{r}'].value).strip() for l, h in cols if ws[f'{l}{r}'].value not in (None, '')}
    toc.append(dict(row_number=r, day_label=str(week).strip(), topic=str(topic).strip(), data=data))
for t in toc: print(t['row_number'], t['day_label'], '|', t['data']['Module'], '|', t['topic'][:70])
TRACK = 'GET 2026 · Java FSD (Track A)'
ms = [('1', 'Core Java & OOP: Community Library Lending System', 2, [10]),
      ('2', 'Design Patterns & Architecture (Python): Notification Dispatcher', 3, [21]),
      ('3', 'Spring Boot Microservices: Secure Task API', 4, [25]),
      ('4', 'Testing, Git & CI/CD: Unit-Testing an Order Service (JUnit 5 + Mockito)', 5, [30]),
      ('6', 'Concurrency, Kafka & Data Engineering: A Concurrent ETL Pipeline', 6, [39])]
q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: '$n$' + s + '$n$'
out = ['begin;',
 f"insert into tracks(client_id,name) select id,{t(TRACK)} from clients where name='Sony' on conflict do nothing;",
 f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,day_column,topic_column)
select c.id,tr.id,{t(FILE)},{t(SHEET)},{HR},{q([{'letter': l, 'header': h} for l, h in cols])}::jsonb,'A','D'
from clients c join tracks tr on tr.client_id=c.id and tr.name={t(TRACK)} where c.name='Sony'
on conflict (client_id,file_name,sheet_name) do update set columns=excluded.columns, track_id=excluded.track_id;""",
 f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,x.day_label,x.topic,x.data from toc_files f, jsonb_to_recordset({q(toc)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
where f.file_name={t(FILE)} and f.sheet_name={t(SHEET)}
on conflict (toc_file_id,row_number) do update set day_label=excluded.day_label, topic=excluded.topic, data=excluded.data;"""]
for seq, name, r, rows in ms:
    src = q({'source': {'file': 'Details_SkillAssist.xlsx', 'sheet': 'Sony GET 2026', 'row': r}})
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(TRACK)} join toc_files f on f.client_id=c.id and f.file_name={t(FILE)} and f.sheet_name={t(SHEET)} where c.name='Sony'),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'Milestone Assessment','{seq}',{t(name)},'2026-07-23',{src}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number = any(array{rows}::int[]) on conflict do nothing;""")
out.append('commit;')
open(base.replace('Skill Assist Tracker/', '') and '/tmp/claude-1000/-home-administrator-Downloads-IntelliJ-Projects-Skill-Assist-Tracker/a76b5546-0d98-436c-b2f1-1898ab9b6cc8/scratchpad/get.sql', 'w').write('\n'.join(out))
