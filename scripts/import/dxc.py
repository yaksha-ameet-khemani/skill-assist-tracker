import openpyxl, json, sys, warnings
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
wb = openpyxl.load_workbook(base + 'DXC/Assessment_Topics.xlsx', data_only=True)
ws = wb['Assessment Topics']
HR = 4
grid = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
for m in ws.merged_cells.ranges:          # fill vertical merges (area / sub-category)
    if m.min_col == m.max_col:
        v = grid.get((m.min_row, m.min_col))
        for r in range(m.min_row + 1, m.max_row + 1): grid[(r, m.min_col)] = v
# area / sub-category also apply to following rows until the next value (not always merged)
area = sub = None
cols = [(c.column_letter, str(c.value).strip()) for c in ws[HR] if c.value is not None]
toc = []
for r in range(HR + 1, ws.max_row + 1):
    topic = grid.get((r, 4))
    if not topic: continue
    area = grid.get((r, 2)) or area
    s = grid.get((r, 3))
    sub = s if s is not None else sub
    if grid.get((r, 2)): sub = s
    data = {'Assessment Area': str(area).strip(), 'Topic': str(topic).strip()}
    if sub and str(sub).strip() != '-': data['Sub-Category'] = str(sub).strip()
    toc.append(dict(row_number=r, day_label=None, topic=str(topic).strip(), data=data, area=str(area).strip()))
by_row = {t['row_number']: t for t in toc}

# content -> (area/track, topic rows)
items = [
 ('.NET', 'DotNet - Order Processing Engine', 2, 'Assessments', [6, 7, 8, 9, 11]),
 ('.NET', 'DotNet - Generic In-Memory Repository', 3, 'Assessments', [7, 8, 9]),
 ('.NET', 'DotNet - Async Retry and Bounded Parallelism', 4, 'Assessments', [10, 11, 12]),
 ('Azure', 'Azure - Secure Storage Account with Bicep', 5, 'Assignments', []),
 ('Azure', 'Azure - HTTP-Triggered Azure Function for Order Intake', 6, 'Assignments', [22]),
 ('Java 8 to Java 25', 'Java - Employee Analytics with Streams', 7, 'Assessments', [38, 39, 41, 48]),
 ('Java 8 to Java 25', 'Java - Library Management System', 8, 'Assessments', [37, 41, 42]),
 ('Java 8 to Java 25', 'Java - Shape Model with Sealed Types and Pattern Matching', 9, 'Assessments', [45, 46]),
 ('GitHub Copilot Fundamentals', 'GitHub - Copilot-Assisted Implementation with Guided Prompts and Review', 10, 'Assignments', [52, 54, 56, 58]),
 ('Prompt Engineering', 'Prompt Engineering - Structured-Output Support Ticket Classifier Prompt', 11, 'Assignments', [63, 64, 67, 68]),
 ('React JS', 'React - Task Manager Component', 12, 'Assessments', [77, 78, 79, 80, 81, 85]),
 ('React JS', 'React - Shopping Cart with useReducer and a Custom Hook', 13, 'Assessments', [79, 83, 84]),
 ('AWS', 'AWS - Secure S3 Bucket and Least-Privilege IAM Policy with Terraform', 14, 'Assignments', [95, 96, 111, 112]),
 ('SQL / SQL Server', 'SQL - Joins and Aggregation Queries', 15, 'Assessments', [118, 119]),
 ('SQL / SQL Server', 'SQL - Window Functions and Common Table Expressions', 16, 'Assessments', [121, 122]),
]
for area_, name, r, created, rows in items:
    for x in rows: assert by_row[x]['area'] == area_, (name, x, by_row[x]['area'])
    print(f"{name[:72]:72} -> " + ('; '.join(f"{x}: {by_row[x]['topic'][:40]}" for x in rows) or 'NO MATCHING TOPIC'), file=sys.stderr)
print(f"{len(toc)} topics, areas: {sorted(set(t['area'] for t in toc))}", file=sys.stderr)

q = lambda o: '$j$' + json.dumps(o) + '$j$'
t = lambda s: '$n$' + s + '$n$'
F = 'Assessment_Topics.xlsx'
out = ['begin;', "insert into clients(name) values ('DXC') on conflict do nothing;"]
for a in dict.fromkeys(i[0] for i in items):
    out.append(f"insert into tracks(client_id,name) select id,{t(a)} from clients where name='DXC' on conflict do nothing;")
out.append(f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,day_column,topic_column)
select id,null,'{F}','Assessment Topics',{HR},{q([{'letter': 'B', 'header': 'Assessment Area'}, {'letter': 'C', 'header': 'Sub-Category'}, {'letter': 'D', 'header': 'Topic'}])}::jsonb,null,'D'
from clients where name='DXC' on conflict (client_id,file_name,sheet_name) do update set columns=excluded.columns;""")
out.append(f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id, x.row_number, null, x.topic, x.data
from toc_files f join clients c on c.id=f.client_id and c.name='DXC', jsonb_to_recordset({q([{k: v for k, v in tt.items() if k != 'area'} for tt in toc])}::jsonb) as x(row_number int, topic text, data jsonb)
where f.file_name='{F}'
on conflict (toc_file_id,row_number) do update set topic=excluded.topic, data=excluded.data;""")
for area_, name, r, created, rows in items:
    extra = q({'created_in': created, 'source': {'file': 'Details_SkillAssist.xlsx', 'sheet': 'DXC CES', 'row': r}})
    out.append(f"""with ctx as (select c.id cid, t.id tid, f.id fid from clients c join tracks t on t.client_id=c.id and t.name={t(area_)} join toc_files f on f.client_id=c.id and f.file_name='{F}' where c.name='DXC'),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'Assessment',null,{t(name)},'2026-08-25',{extra}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,tr.id from ins cross join ctx join toc_rows tr on tr.toc_file_id=ctx.fid and tr.row_number = any(array{rows or [0]}::int[]) on conflict do nothing;""")
out.append('commit;')
print('\n'.join(out))
