# BCT Banking sheet of Details_SkillAssist.xlsx (user-confirmed 2026-10-03): CSM Ayman.
# Phase 1 -> track "Pre-Learning", all 4 items Milestone Assessment (day blank). Phase 3 & 4 – FSE and Phase – 3 keep
# the sheet's names: assignments = Daily Assignment numbered in row order per track, linked to their Topic row;
# "Milestone N" rows = Milestone Assessment linked to the topics since the previous milestone.
# The sheet's Module + Topic columns are the TOC. Dates from the merged Delivery Date column.
import openpyxl, json, re, warnings, sys
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
FILE, SHEET, CLIENT, CSM = 'Details_SkillAssist.xlsx', 'BCT Banking', 'BCT Banking', 'Ayman'
RENAME = {'Phase – 1 Pre-Learning': 'Pre-Learning'}
one = lambda v: None if v is None or not str(v).strip() else re.sub(r'\s+', ' ', str(v)).strip()
iso = lambda d: '-'.join(reversed(d.split('-')))
ws = openpyxl.load_workbook(base + FILE, data_only=True)[SHEET]

toc, items = [], []
track = module = date = topic_row = None
for r in range(2, ws.max_row + 1):
    name = one(ws[f'E{r}'].value)
    if not name: continue
    if one(ws[f'A{r}'].value):
        sheet_track = one(ws[f'A{r}'].value)
        track, module, date, day, since_m = RENAME.get(sheet_track, sheet_track), one(ws[f'B{r}'].value), iso(ws[f'F{r}'].value), 0, []
    c = one(ws[f'C{r}'].value)
    m = re.match(r'Milestone (\d+)$', c or '')
    if c and not m:
        topic_row = r; since_m.append(r)
        toc.append(dict(row_number=r, day_label=None, topic=c, data={'Track': sheet_track, 'Module': module, 'Topic': c}))
    if m:
        items.append(dict(row=r, track=track, typ='Milestone Assessment', seq=m.group(1), name=name, date=date, rows=since_m))
        since_m = []
    elif track == 'Pre-Learning':
        items.append(dict(row=r, track=track, typ='Milestone Assessment', seq=None, name=name, date=date, rows=[topic_row]))
    else:
        day += 1
        items.append(dict(row=r, track=track, typ='Daily Assignment', seq=str(day), name=name, date=date, rows=[topic_row]))
by_row = {t['row_number']: t['topic'] for t in toc}
for i in items: print(f"{i['row']:>2} {i['track'][:14]:14} {i['typ'][:9]} {i['seq'] or '-':>2} {i['date']} {i['name'][:50]:50} -> {[by_row[x] for x in i['rows']]}", file=sys.stderr)
print(len(toc), 'topics;', len(items), 'items', file=sys.stderr)

q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
out = ['begin;', f"insert into clients(name) values ({t(CLIENT)}) on conflict do nothing;"]
for tr in dict.fromkeys(i['track'] for i in items):
    out.append(f"insert into tracks(client_id,name,csm) select id,{t(tr)},{t(CSM)} from clients where name={t(CLIENT)} on conflict (client_id,name) do update set csm=excluded.csm;")
cols = [{'letter': 'A', 'header': 'Track'}, {'letter': 'B', 'header': 'Module'}, {'letter': 'C', 'header': 'Topic'}]
out.append(f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,topic_column)
select id,null,{t(FILE)},{t(SHEET)},1,{q(cols)}::jsonb,'C' from clients where name={t(CLIENT)}
on conflict (client_id,file_name,sheet_name) do update set columns=excluded.columns;""")
out.append(f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,x.day_label,x.topic,x.data from toc_files f join clients c on c.id=f.client_id and c.name={t(CLIENT)}, jsonb_to_recordset({q(toc)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
where f.file_name={t(FILE)} and f.sheet_name={t(SHEET)}
on conflict (toc_file_id,row_number) do update set topic=excluded.topic, data=excluded.data;""")
for i in items:
    extra = q({'source': {'file': FILE, 'sheet': SHEET, 'row': i['row']}})
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(i['track'])} join toc_files f on f.client_id=c.id and f.file_name={t(FILE)} and f.sheet_name={t(SHEET)} where c.name={t(CLIENT)}),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'{i['typ']}',{t(i['seq'])},{t(i['name'])},'{i['date']}',{extra}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number = any(array{i['rows']}::int[]) on conflict do nothing;""")
out.append('commit;')
print('\n'.join(out))
