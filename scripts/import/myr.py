import openpyxl, json, re, warnings, sys
warnings.filterwarnings('ignore')
ws = openpyxl.load_workbook('/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/Details_SkillAssist.xlsx', data_only=True)['Myridious 2026']
def date(r): return '2026-07-20' if r <= 11 else '2026-07-22' if r <= 21 else '2026-07-23'
clean = lambda v: re.sub(r'\s+', ' ', str(v)).strip() if v is not None else ''
items = []
for r in range(2, ws.max_row + 1):
    lp, course, an, q, skill, name = (clean(ws.cell(r, c).value) for c in range(1, 7))
    if not name: continue
    m = re.search(r'milestone(?: assessment)?\s*(\d+)', course, re.I)
    typ, seq = ('Milestone Assessment', m.group(1)) if m else ('Assignment', None)
    items.append(dict(row=r, track=lp or 'No LP Name', typ=typ, seq=seq, name=name, skill=skill, date=date(r),
                      extra={'assessment_name': an, 'course_name': course or None,
                             'source': {'file': 'Details_SkillAssist.xlsx', 'sheet': 'Myridious 2026', 'row': r}}))
for i in items: print(f"{i['row']:>2} {str(i['track']):32} {i['typ'][:9]:9} {i['seq'] or '':2} {i['name'][:55]:55} [{i['skill']}] {i['date']}", file=sys.stderr)
skills = sorted({i['skill'] for i in items})
print(len(items), 'items;', len(skills), 'skills:', skills, file=sys.stderr)
q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
out = ['begin;',
 "insert into content_types(name, sort_order) values ('Assignment', 52) on conflict do nothing;",
 "insert into clients(name) values ('Myridius') on conflict do nothing;"]
for tr in sorted({i['track'] for i in items if i['track']}):
    out.append(f"insert into tracks(client_id,name,csm) select id,{t(tr)},'Internal Request' from clients where name='Myridius' on conflict (client_id,name) do update set csm=excluded.csm;")
out.append("""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,day_column,topic_column,notes)
select id,null,'(Added manually)','Myridius skills',1,'[{"letter":"A","header":"Topic"}]'::jsonb,null,null,'No TOC file for Myridius; each item''s skill_name used as its topic.'
from clients where name='Myridius' on conflict (client_id,file_name,sheet_name) do nothing;""")
for n, sk in enumerate(skills, start=2):
    out.append(f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,{n},null,{t(sk)},{q({'Topic': sk})}::jsonb from toc_files f join clients c on c.id=f.client_id and c.name='Myridius' where f.file_name='(Added manually)'
on conflict (toc_file_id,row_number) do update set topic=excluded.topic, data=excluded.data;""")
for i in items:
    trk = f"(select tr.id from tracks tr where tr.client_id=c.id and tr.name={t(i['track'])})" if i['track'] else 'null'
    out.append(f"""with ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select c.id,{trk},'{i['typ']}',{t(i['seq'])},{t(i['name'])},'{i['date']}',{q(i['extra'])}::jsonb from clients c where c.name='Myridius'
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id, r.id from ins, toc_rows r join toc_files f on f.id=r.toc_file_id join clients c on c.id=f.client_id and c.name='Myridius'
where f.file_name='(Added manually)' and r.topic={t(i['skill'])} on conflict do nothing;""")
out.append('commit;')
open('/tmp/claude-1000/-home-administrator-Downloads-IntelliJ-Projects-Skill-Assist-Tracker/a76b5546-0d98-436c-b2f1-1898ab9b6cc8/scratchpad/myr.sql', 'w').write('\n'.join(out))
