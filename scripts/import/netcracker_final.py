# Netcracker sheet of Details_SkillAssist.xlsx, rows 64-66 (user-confirmed 2026-10-04): "Final Assessment" of
# "Netcracker Modular Learning 20 Feb 2026" -> track "Final - Java", CSM Aarti, all Milestone Assessment (no number),
# 18-03-2026. Column B -> Course column, column A kept as extra.program. TOC typed in from the screenshot shared on
# 2026-10-04 (topic = On-the-Ground Developer Skill); each final assessment links to every topic, as milestones do.
import openpyxl, json, re, warnings, sys
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
FILE, SHEET, CLIENT, CSM, TRACK, DATE = 'Details_SkillAssist.xlsx', 'Netcracker', 'Netcracker', 'Aarti', 'Final - Java', '2026-03-18'
FN, SH = '(Added manually) Netcracker TOC screenshot', 'Final - Java'
TOC = [  # (skill, knowledge / learning item, course)
    ('Logic & Control Flow', 'Writing clean, non-nested conditional logic; effective use of Switch expressions (Java 17+).', 'Java from Beginner to Expert - Second Edition'),
    ('Memory-Efficient String Handling', 'Using StringBuilder/StringBuffer for heavy manipulation; understanding the String Constant Pool.', 'Java 21 Programming Masterclass: Fundamentals for Beginners'),
    ('Type Safety & Generics', 'Creating generic classes and methods to ensure type safety across product modules.', 'Java from Beginner to Expert - Second Edition'),
    ('Boilerplate Reduction', 'Using Java Records (Java 14+) for immutable data carriers (DTOs) instead of verbose POJOs.', None),
    ('Interface-Based Programming', 'Designing decoupled systems using Interfaces; understanding "Default" and "Static" interface methods.', 'Modern Java - Mastering Features from Java 8 to Java 21'),
    ('Inheritance & Sealed Classes', 'Using extends for hierarchy and Sealed Classes (Java 17) to restrict which classes can implement a contract.', 'Modern Java - Mastering Features from Java 8 to Java 21; Mastering Advanced Java with Object-Oriented Programming'),
    ('Encapsulation & Immutability', 'Protecting business logic by designing strictly immutable objects and using private visibility correctly.', 'Mastering Advanced Java with Object-Oriented Programming'),
    ('Functional Programming', 'Implementing Lambdas and the Streams API to replace legacy "For-loops" for data processing.', 'Modern Java - Mastering Features from Java 8 to Java 21'),
    ('Collection Selection', 'Choosing the right structure (e.g., ArrayList vs LinkedList, HashMap vs TreeMap) based on Big-O performance.', 'Java Mastery Intermediate: Methods, Collections, and Beyond'),
    ('Thread-Safe Data Handling', 'Implementing ConcurrentHashMap and CopyOnWriteArrayList in multi-user environments.', 'Java Concurrency and Multithreading in Practice'),
    ('Stream Pipelines', 'Writing complex .filter(), .map(), and .collect() chains with proper Optional handling to avoid NullPointerExceptions.', None),
    ('Sequenced Collections', 'Utilizing the new Java 21 SequencedCollection interfaces for predictable element ordering.', 'Java 21 - Exploring the Latest Innovations for 2024'),
    ('High-Scale Concurrency', 'Implementing Virtual Threads (Java 21/Project Loom) to handle thousands of concurrent API requests.', 'Java 21 - Exploring the Latest Innovations for 2024')]
ws = openpyxl.load_workbook(base + FILE, data_only=True)[SHEET]
one = lambda v: re.sub(r'\s+', ' ', str(v)).strip()
assert one(ws['C64'].value) == 'Final Assessment' and one(ws['F64'].value) == '18-03-2026'
program, course = one(ws['A64'].value), one(ws['B64'].value)
items = [dict(row=r, name=one(ws[f'E{r}'].value)) for r in (64, 65, 66)]
toc = [dict(row_number=n, day_label=None, topic=sk, data={k: v for k, v in
            {'On-the-Ground Developer Skill': sk, 'Knowledge / Learning Item': ki, 'Course': co}.items() if v})
       for n, (sk, ki, co) in enumerate(TOC, 2)]
for i in items: print(i['row'], i['name'], '-> all', len(toc), 'topics', file=sys.stderr)

q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
cols = [{'letter': 'A', 'header': 'On-the-Ground Developer Skill'}, {'letter': 'B', 'header': 'Knowledge / Learning Item'}, {'letter': 'C', 'header': 'Course'}]
note = ('Typed in from the screenshot shared on 2026-10-04. Some text is red in the screenshot (Logic & Control Flow, String Constant Pool, '
        'Boilerplate Reduction). The screenshot is cut off after High-Scale Concurrency (a heap-dump / JProfiler row is only partly visible and is not loaded).')
out = ['begin;',
       f"insert into tracks(client_id,name,csm) select id,{t(TRACK)},{t(CSM)} from clients where name={t(CLIENT)} on conflict (client_id,name) do update set csm=excluded.csm;",
       f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,topic_column,notes)
select c.id,(select id from tracks where client_id=c.id and name={t(TRACK)}),{t(FN)},{t(SH)},1,{q(cols)}::jsonb,'A',{t(note)}
from clients c where c.name={t(CLIENT)} on conflict (client_id,file_name,sheet_name) do update set notes=excluded.notes;""",
       f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,x.day_label,x.topic,x.data from toc_files f join clients c on c.id=f.client_id and c.name={t(CLIENT)}, jsonb_to_recordset({q(toc)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
where f.file_name={t(FN)} and f.sheet_name={t(SH)}
on conflict (toc_file_id,row_number) do update set topic=excluded.topic, data=excluded.data;"""]
for i in items:
    extra = q({'source': {'file': FILE, 'sheet': SHEET, 'row': i['row']}, 'course': course, 'program': program})
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(TRACK)} join toc_files f on f.client_id=c.id and f.file_name={t(FN)} and f.sheet_name={t(SH)} where c.name={t(CLIENT)}),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'Milestone Assessment',null,{t(i['name'])},'{DATE}',{extra}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid on conflict do nothing;""")
out.append('commit;')
print('\n'.join(out))
