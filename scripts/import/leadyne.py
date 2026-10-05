# Leadyne sheet of Details_SkillAssist.xlsx, all rows 2-16 (user-confirmed 2026-10-04): CSM Divyabharthi; date from column E
# (30-01-2026 for rows 2-11, 18-02-2026 for the Junit block 13-16).
# Track = column A (Module). Assignment -> Daily Assignment; "Milestone 1" -> Milestone Assessment 1 (TOC: Module-End
# Assessment); "Incremental Project" -> Capstone with extra.project = "Incremental Project N" (home-page Project column; N
# from the TOC or from column B). Column B -> Course column. TOC typed in from three screenshots (topic = TOC "Topic"; for
# JUnit the TOC's "Coverage Intent" rows are the topics). A daily assignment links to its topic, or for "After <topic>" to
# the topics since the previous assignment up to that one; a milestone / incremental project to every topic of its module.
import openpyxl, json, re, warnings, sys
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
FILE, SHEET, CLIENT, CSM = 'Details_SkillAssist.xlsx', 'Leadyne', 'Leadyne', 'Divyabharthi'
FN, SH = '(Added manually) Leadyne TOC screenshot', 'Leadyne TOC'
EP, TA, JU = 'Engineering & Programming Foundations', 'Test Automation, CI/CD & AI-Driven QE', 'Junit'
TOC = [  # (module, track, topic, coverage intent, sub-topic, self-paced hrs, instructor-led hrs)
    (EP, 'Agile & Delivery', 'Agile Delivery & Agile Practice Group', 'Agile execution in BFSI enterprises', 'Agile ceremonies, BFSI delivery, shift-left quality', '1', '0'),
    (EP, 'Design & Engineering Fundamentals', 'Object-Oriented Programming (OOP from Scratch)', 'From basics to enterprise object modeling', 'OOP principles, object modeling, classes, abstraction, inheritance, polymorphism', '4', '2'),
    (EP, 'Design & Engineering Fundamentals', 'Software Design Patterns (Basic to Advanced)', 'Core + enterprise usage patterns', 'SOLID principles, creational & structural patterns, enterprise usage', '4', '2'),
    (EP, 'Design & Engineering Fundamentals', 'Advanced Design Patterns (Deep Dive)', 'Complex, real-world system design', 'Behavioral patterns, scalability patterns, real-world trade-offs', '4', '2'),
    (EP, 'Design & Engineering Fundamentals', 'Clean Code & Refactoring', 'Maintainable, readable, legacy code', 'Clean code practices, refactoring techniques, code smells, legacy improvement', '2', '2'),
    (EP, 'Design & Engineering Fundamentals', 'Design Mindset & System Thinking', 'How senior engineers think', 'System thinking, architectural mindset, decision-making', '1', '1'),
    (EP, 'Programming Foundations', 'Java Programming (Zero to Advanced)', 'Core Java, OOP, collections, concurrency', 'Core Java, OOP, collections, concurrency, multithreading', '6', '6'),
    (EP, 'Programming Foundations', 'Python Programming (Zero to Advanced)', 'Core Python, automation, services', 'Core Python, OOP, automation scripting, service development', '2', '4'),
    (TA, 'Test Management', 'JIRA & Zephyr – Test Cases & Cycles', 'Tool mastery & traceability', 'Test management tools, traceability', '2', '0'),
    (TA, 'Test Management', 'BDD / TDD / OATS', 'Shift-left testing practices', 'BDD, TDD, test-first development', '3', '0'),
    (TA, 'Test Management', 'Test Types (Functional to Non-Functional)', 'Full coverage model', 'Functional, performance, security testing', '0', '1.5'),
    (TA, 'Automation Engineering', 'Functional Test Automation (Zero to Advanced)', 'Frameworks & best practices',
     'Automation frameworks, scripting - Locators and Object identification (By ID, By name, By class name, By Tag name, By Xpath, By CSS); '
     'Handling different controls on web page – Basic (button, input box, checkbox, radio button, select box); Handling different controls on web page – '
     'Advanced (Alert Box, Datepicker, multiple windows/Tabs, drag and drop, Iframes, Dynamic objects)', '5', '0'),
    (TA, 'Automation Engineering', 'Non-Functional Testing', 'Performance, security, reliability', 'Load, security, reliability testing', '3', '0'),
    (TA, 'Automation Engineering', 'Client Automation Frameworks', 'Adapting existing frameworks', 'Framework customization, enterprise adaptation', '3', '0'),
    (TA, 'Automation Engineering', 'Automation CI/CD Integration', 'DevOps-aligned QE', 'CI/CD pipelines, automated quality gates', '4', '0'),
    (TA, 'AI in QE', 'AI-Driven Test Automation', 'Prompt-based automation', 'AI-generated tests, prompts', '4', '0'),
    (TA, 'AI in QE', 'AI QE Agents & Self-Healing Tests', 'Autonomous QE systems', 'Self-healing tests, autonomous agents', '4', '0'),
    # JUnit screenshot: Topic "JUNIT" spans all rows; each Coverage Intent row is a topic; hours (6 / 8) are for the whole block
    (JU, 'JUNIT', 'Testing in BFSI Industry Context', 'Testing in BFSI Industry Context', 'Importance of Testing in BFSI, Financial Impact of Defects, Regulatory & Compliance Requirements, Transaction Scale & Risk Factors', '6', '8'),
    (JU, 'JUNIT', 'JUnit 5 Fundamentals', 'JUnit 5 Fundamentals', 'JUnit Architecture, Jupiter–Platform–Vintage, Test Lifecycle, Unit Test Structure, JUnit Evolution', '6', '8'),
    (JU, 'JUNIT', 'JUnit Annotations Framework', 'JUnit Annotations Framework', 'Core Annotations, Lifecycle Annotations, Display & Nested Tests, Annotation Usage', '6', '8'),
    (JU, 'JUNIT', 'Assertions & Validation', 'Assertions & Validation', 'Assertions Library, Equality & Boolean Assertions, Exception Testing, Group Assertions, Timeout Assertions', '6', '8'),
    (JU, 'JUNIT', 'Test Design Foundations', 'Test Design Foundations', 'AAA Pattern, Test Data Preparation, Execution Strategies, Validation Techniques', '6', '8'),
    (JU, 'JUNIT', 'Code Coverage Concepts', 'Code Coverage Concepts', 'Coverage Metrics, Line vs Branch Coverage, BFSI Coverage Standards, Coverage Interpretation', '6', '8'),
    (JU, 'JUNIT', 'Testing Best Practices', 'Testing Best Practices', 'Naming Conventions, Test Independence, Boundary Testing, Mocking Fundamentals', '6', '8')]
PROJECT = {EP: 'Incremental Project 1', TA: 'Incremental Project 4'}
# sheet column B (course / track) -> TOC topic for daily assignments
DAILY = {('Programming Foundations', 'Java'): 'Java Programming (Zero to Advanced)',
         ('Programming Foundations', 'Python'): 'Python Programming (Zero to Advanced)',
         ('Functional Test Automation (Zero to Advanced)', 'Java Selenium'): 'Functional Test Automation (Zero to Advanced)'}
ws = openpyxl.load_workbook(base + FILE, data_only=True)[SHEET]
one = lambda v: re.sub(r'\s+', ' ', str(v)).strip()
assert one(ws['E2'].value) == '30-01-2026' and one(ws['E13'].value) == '18-02-2026'
toc = [dict(row_number=n, day_label=None, topic=tp, data={'Modules': m, 'Track': tr, 'Topic': tp, 'Coverage Intent': ci, 'Sub-Topic': st,
            'Self-Paced Learning Hours': sp, 'Instructor-Led Learning Hours': il}) for n, (m, tr, tp, ci, st, sp, il) in enumerate(TOC, 2)]
items, mod, col_b, date, since = [], None, None, None, 0
for r in list(range(2, 12)) + list(range(13, 17)):
    if ws[f'A{r}'].value: mod, since = one(ws[f'A{r}'].value), 0
    if ws[f'E{r}'].value: d, m, y = one(ws[f'E{r}'].value).split('-'); date = f'{y}-{m}-{d}'
    col_b = one(ws[f'B{r}'].value) if ws[f'B{r}'].value else col_b
    name = one(ws[f'D{r}'].value)
    whole = [x['row_number'] for x in toc if x['data']['Modules'] == mod]
    extra = {}
    if col_b.lower() in ('milestone 1', 'module end assessment'): typ, seq, hit = 'Milestone Assessment', '1', whole
    elif col_b.startswith('Incremental Project'):
        typ, seq, hit = 'Capstone', None, whole; extra['project'] = col_b if col_b[-1].isdigit() else PROJECT[mod]
    elif col_b.startswith('After '):  # "After <topic>": topics since the previous assignment, up to and including <topic>
        typ, seq = 'Daily Assignment', None
        upto = whole.index(next(x['row_number'] for x in toc if x['data']['Modules'] == mod and x['topic'] == col_b[6:]))
        hit = whole[since:upto + 1]; since = upto + 1
    else:
        typ, seq = 'Daily Assignment', None
        hit = [x['row_number'] for x in toc if x['topic'] == DAILY[(col_b, name.split(' - ')[0])]]
    assert hit, r
    items.append(dict(row=r, track=mod, type=typ, seq=seq, name=name, course=col_b, hit=hit, extra=extra, date=date))
    print(r, mod[:20], typ, seq or '', extra.get('project', ''), name, '->', len(hit), file=sys.stderr)

q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
cols = [{'letter': 'A', 'header': 'Modules'}, {'letter': 'B', 'header': 'Track'}, {'letter': 'C', 'header': 'Topic'}, {'letter': 'D', 'header': 'Coverage Intent'},
        {'letter': 'E', 'header': 'Sub-Topic'}, {'letter': 'F', 'header': 'Self-Paced Learning Hours'}, {'letter': 'G', 'header': 'Instructor-Led Learning Hours'}]
note = ('Typed in from three screenshots shared on 2026-10-04. Module-End Assessment / Incremental Project rows are markers, not topics. '
        'The second screenshot starts mid-sheet (a row above JIRA & Zephyr is cut off) and the Functional Test Automation sub-topic list is '
        'partly cut off at top and bottom; only visible text is kept. Hours shown merged across rows are repeated on each row. '
        'JUnit: the TOC Topic is "JUNIT" for the whole block, so its Coverage Intent rows are used as topics.')
out = ['begin;', f"insert into clients(name) values ({t(CLIENT)}) on conflict do nothing;"]
for tr in dict.fromkeys(i['track'] for i in items):
    out.append(f"insert into tracks(client_id,name,csm) select id,{t(tr)},{t(CSM)} from clients where name={t(CLIENT)} on conflict (client_id,name) do update set csm=excluded.csm;")
out += [f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,topic_column,notes)
select c.id,null,{t(FN)},{t(SH)},1,{q(cols)}::jsonb,'C',{t(note)} from clients c where c.name={t(CLIENT)}
on conflict (client_id,file_name,sheet_name) do update set notes=excluded.notes;""",
        f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,x.day_label,x.topic,x.data from toc_files f join clients c on c.id=f.client_id and c.name={t(CLIENT)}, jsonb_to_recordset({q(toc)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
where f.file_name={t(FN)} and f.sheet_name={t(SH)}
on conflict (toc_file_id,row_number) do update set topic=excluded.topic, data=excluded.data;"""]
for i in items:
    extra = q({'source': {'file': FILE, 'sheet': SHEET, 'row': i['row']}, 'course': i['course'], **i['extra']})
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(i['track'])} join toc_files f on f.client_id=c.id and f.file_name={t(FN)} and f.sheet_name={t(SH)} where c.name={t(CLIENT)}),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,{t(i['type'])},{t(i['seq'])},{t(i['name'])},'{i['date']}',{extra}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number = any(array{i['hit']}::int[]) on conflict do nothing;""")
out.append('commit;')
print('\n'.join(out))
