# EY GDS sheet of Details_SkillAssist.xlsx, rows 58-105 "AI & Azure" (user-confirmed 2026-10-03): track AI & Azure,
# CSM Muzzamil, all dated 29-09-2026. Content = "Questions" column J. ShopEase (row 59) is the Pre-Assessment;
# "Milestone N" rows are Milestone Assessments linked to that module's topics since the previous milestone
# (Capstone days excluded); every other question is a practice assignment stored as Daily Assignment on its day,
# linked to all topics above it back to the previous question.
# TOC = EY GDS/AI & Azure-Day Wise-Planner & Assessment Schedule (1).xlsx, "40-Day Planner": topic = Module / Topic,
# subtopics = Detailed Coverage; its Wk column (W1-W8) is kept as "Week N" on topics and items (Week column). Its rows are the content sheet's rows minus 56.
import openpyxl, json, re, warnings, sys
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
TFILE, TSHEET, OFF = 'AI & Azure-Day Wise-Planner & Assessment Schedule (1).xlsx', '40-Day Planner', 56
FILE, SHEET, CLIENT, TRACK, CSM, DATE = 'Details_SkillAssist.xlsx', 'EY GDS', 'EY GDS', 'AI & Azure', 'Muzzamil', '2026-09-29'
one = lambda v: None if v is None or not str(v).strip() else re.sub(r'\s+', ' ', str(v)).strip()

tw = openpyxl.load_workbook(base + 'EY GDS/' + TFILE, data_only=True)[TSHEET]
ws = openpyxl.load_workbook(base + FILE, data_only=True)[SHEET]
toc, track_of, week_of, day, wk = [], {}, {}, None, None
for r in range(3, tw.max_row + 1):
    if tw[f'A{r}'].value: day = str(tw[f'A{r}'].value)
    if tw[f'C{r}'].value: wk = 'Week ' + re.sub(r'\D', '', str(tw[f'C{r}'].value))   # "W3" -> "Week 3"
    week_of[r] = wk
    topic = one(tw[f'E{r}'].value)
    assert topic == one(ws[f'E{r + OFF}'].value), f'row {r} no longer lines up'
    if not topic: continue
    toc.append(dict(row_number=r, day_label=day, topic=topic,
                    data={'Week': wk, 'Day': day, 'Module / Topic': topic, 'Detailed Coverage': one(tw[f'F{r}'].value)}))
    track_of[r] = one(tw[f'D{r}'].value)
by_row = {t['row_number']: t for t in toc}

items, since_q, since_m, day = [], [], [], None
for r in range(59, ws.max_row + 1):
    tr = r - OFF
    if ws[f'A{r}'].value: day = str(ws[f'A{r}'].value)
    if tr in by_row: since_q.append(tr); since_m.append(tr)
    name, label = one(ws[f'J{r}'].value), one(ws[f'D{r}'].value) or ''
    if not name: continue
    m = re.match(r'Milestone (\d+)', label)
    if m:
        items.append(dict(row=r, typ='Milestone Assessment', seq=m.group(1), name=name,
                          rows=[x for x in since_m if track_of[x] != 'Capstone']))
        since_q, since_m = [], []
    elif name.startswith('ShopEase: Mini'):
        items.append(dict(row=r, typ='Pre-Assessment', seq=day, name=name, rows=[tr]))
    else:
        items.append(dict(row=r, typ='Daily Assignment', seq=day, name=name, rows=list(since_q)))
        since_q = []
for i in items: print(f"{i['row']:>3} {i['typ'][:10]:10} {i['seq']:>2} {i['name'][:50]:50} -> D{'/D'.join(dict.fromkeys(by_row[x]['day_label'] for x in i['rows']))}", file=sys.stderr)
print(len(toc), 'TOC rows;', len(items), 'items', file=sys.stderr)

q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
out = ['begin;', f"insert into clients(name) values ({t(CLIENT)}) on conflict do nothing;",
       f"insert into tracks(client_id,name,csm) select id,{t(TRACK)},{t(CSM)} from clients where name={t(CLIENT)} on conflict (client_id,name) do update set csm=excluded.csm;",
       f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,day_column,topic_column)
select c.id,tr.id,{t(TFILE)},{t(TSHEET)},2,{q([{'letter': 'C', 'header': 'Week'}, {'letter': 'A', 'header': 'Day'}, {'letter': 'E', 'header': 'Module / Topic'}, {'letter': 'F', 'header': 'Detailed Coverage'}])}::jsonb,'A','E' from clients c join tracks tr on tr.client_id=c.id and tr.name={t(TRACK)} where c.name={t(CLIENT)}
on conflict (client_id,file_name,sheet_name) do update set columns=excluded.columns;""",
       f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,x.day_label,x.topic,x.data from toc_files f join clients c on c.id=f.client_id and c.name={t(CLIENT)}, jsonb_to_recordset({q(toc)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
where f.file_name={t(TFILE)} and f.sheet_name={t(TSHEET)}
on conflict (toc_file_id,row_number) do update set day_label=excluded.day_label, topic=excluded.topic, data=excluded.data;"""]
for i in items:
    extra = q({'week': week_of[i['row'] - OFF], 'source': {'file': FILE, 'sheet': SHEET, 'row': i['row']}})
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(TRACK)} join toc_files f on f.client_id=c.id and f.file_name={t(TFILE)} and f.sheet_name={t(TSHEET)} where c.name={t(CLIENT)}),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'{i['typ']}',{t(i['seq'])},{t(i['name'])},'{DATE}',{extra}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number = any(array{i['rows']}::int[]) on conflict do nothing;""")
out.append('commit;')
print('\n'.join(out))
