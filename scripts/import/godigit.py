# GoDigit sheet of Details_SkillAssist.xlsx (user-approved 2026-10-03): CSM Ayman, tracks = Skill column.
# The sheet's "Focus / Topics" column is the TOC. Assignments = Daily Assignment numbered Day 1.. in row order per
# track, linked to their (merged) topic row; "Milestone N" rows = Milestone Assessment linked to all topics of the track.
# Date = unlabelled column E (06-04-2026, later additions) when present, else column D (09-03-2026).
import openpyxl, json, re, warnings, sys
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
FILE, SHEET, CLIENT, CSM = 'Details_SkillAssist.xlsx', 'GoDigit', 'GoDigit', 'Ayman'
one = lambda v: None if v is None or not str(v).strip() else re.sub(r'\s+', ' ', str(v)).strip()
iso = lambda d: '-'.join(reversed(d.split('-')))
ws = openpyxl.load_workbook(base + FILE, data_only=True)[SHEET]

toc, items, track, trows, topic_row, ddate, day = [], [], None, [], None, None, 0
for r in range(2, ws.max_row + 1):
    name = one(ws[f'C{r}'].value)
    if one(ws[f'A{r}'].value): track, trows, day = one(ws[f'A{r}'].value), [], 0
    if ws[f'D{r}'].value: ddate = iso(ws[f'D{r}'].value)
    b = one(ws[f'B{r}'].value)
    m = re.match(r'Milestone (\d+)$', b or '')
    if b and not m:
        topic_row = r; trows.append(r)
        toc.append(dict(row_number=r, day_label=None, topic=b, data={'Skill': track, 'Focus / Topics': b}))
    if not name: continue
    date = iso(ws[f'E{r}'].value) if ws[f'E{r}'].value else ddate
    if m:
        items.append(dict(row=r, track=track, typ='Milestone Assessment', seq=m.group(1), name=name, date=date, rows=list(trows)))
    else:
        day += 1
        items.append(dict(row=r, track=track, typ='Daily Assignment', seq=str(day), name=name, date=date, rows=[topic_row]))
by_row = {t['row_number']: t['topic'] for t in toc}
for i in items: print(f"{i['row']:>2} {i['track'][:12]} {i['typ'][:9]} {i['seq']:>2} {i['date']} {i['name'][:55]:55} -> {[by_row[x][:25] for x in i['rows']]}", file=sys.stderr)
print(len(toc), 'topics;', len(items), 'items', file=sys.stderr)

q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
out = ['begin;', f"insert into clients(name) values ({t(CLIENT)}) on conflict do nothing;"]
for tr in dict.fromkeys(i['track'] for i in items):
    out.append(f"insert into tracks(client_id,name,csm) select id,{t(tr)},{t(CSM)} from clients where name={t(CLIENT)} on conflict (client_id,name) do update set csm=excluded.csm;")
cols = [{'letter': 'A', 'header': 'Skill'}, {'letter': 'B', 'header': 'Focus / Topics'}]
out.append(f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,topic_column)
select id,null,{t(FILE)},{t(SHEET)},1,{q(cols)}::jsonb,'B' from clients where name={t(CLIENT)}
on conflict (client_id,file_name,sheet_name) do update set columns=excluded.columns;""")
out.append(f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,x.day_label,x.topic,x.data from toc_files f join clients c on c.id=f.client_id and c.name={t(CLIENT)}, jsonb_to_recordset({q(toc)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
where f.file_name={t(FILE)} and f.sheet_name={t(SHEET)}
on conflict (toc_file_id,row_number) do update set topic=excluded.topic, data=excluded.data;""")
for i in items:
    extra = q({'source': {'file': FILE, 'sheet': SHEET, 'row': i['row']}})
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(i['track'])} join toc_files f on f.client_id=c.id and f.file_name={t(FILE)} and f.sheet_name={t(SHEET)} where c.name={t(CLIENT)}),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'{i['typ']}',{t(i['seq'])},{t(i['name'])},'{i['date']}',{extra}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number = any(array{i['rows']}::int[]) on conflict do nothing;""")
out.append('commit;')
print('\n'.join(out))
