"""Build SQL to load one Sony track: TOC rows + daily assignments + milestones + pre-assessment.
Usage: python3 load_track.py config.json  -> prints match preview to stderr, SQL to stdout."""
import openpyxl, re, difflib, json, sys, warnings
warnings.filterwarnings('ignore')
cfg = json.load(open(sys.argv[1]))
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
ws = openpyxl.load_workbook(base + cfg['toc_path'], data_only=True)[cfg['toc_sheet']]
HR, DC, TC = cfg['header_row'], cfg['day_col'], cfg['topic_col']
cols = [(c.column_letter, str(c.value).strip()) for c in ws[HR] if c.value is not None]
toc = []
for r in range(HR + 1, ws.max_row + 1):
    topic = ws[f'{TC}{r}'].value
    if not topic: continue
    data = {h: str(ws[f'{l}{r}'].value).strip() for l, h in cols if ws[f'{l}{r}'].value not in (None, '')}
    toc.append(dict(row_number=r, day_label=str(ws[f'{DC}{r}'].value or '').strip() or None, topic=str(topic).strip(), data=data))

def dnum(s):
    m = re.search(r'day\s*(\d+)', s or '', re.I)
    if m: return int(m.group(1))
    m = re.search(r'\d+', s or '')
    return int(m.group(0)) if m else None

pr = openpyxl.load_workbook(base + 'Details_SkillAssist.xlsx', data_only=True)['Sony Priyanka']
d = cfg['daily']
items = []
for r in range(d['first_row'], d['last_row'] + 1):
    name = pr[f"{d['name_col']}{r}"].value
    if not name: continue
    seq = str(pr[f"{d['day_col']}{r}"].value).strip()
    date = next(dt for upto, dt in d['dates'] if r <= upto)
    cands = [t for t in toc if dnum(t['day_label']) == dnum(seq)]
    best = max(cands, key=lambda t: difflib.SequenceMatcher(None, name.lower(), t['topic'].lower()).ratio()) if cands else None
    items.append(('Daily Assignment', seq, str(name).strip(), date, r, [best['row_number']] if best else []))
    print(f"{seq:>5} | {str(name)[:68]:68} -> {('row %d: %s' % (best['row_number'], best['topic'][:55])) if best else 'NO MATCH'}"
          f"{'  (of %d)' % len(cands) if len(cands) > 1 else ''}", file=sys.stderr)
for m in cfg['milestones']:
    items.append(('Milestone Assessment', m['seq'], m['name'], m['date'], m['sheet_row'], list(range(m['from'], m['to'] + 1))))
for p in cfg.get('pre', []):
    items.append(('Pre-Assessment', None, p['name'], p['date'], p['sheet_row'], p['rows']))
print(f"{len(toc)} TOC rows, {len(items)} content items", file=sys.stderr)

q = lambda o: '$j$' + json.dumps(o) + '$j$'
t = lambda s: '$n$' + s + '$n$'
ctx = (f"select c.id cid, t.id tid, f.id fid from clients c join tracks t on t.client_id=c.id "
       f"join toc_files f on f.track_id=t.id and f.file_name={t(cfg['toc_file_name'])} "
       f"where c.name={t(cfg['client'])} and t.name={t(cfg['track'])}")
out = ['begin;',
       f"insert into clients(name) values ({t(cfg['client'])}) on conflict do nothing;",
       f"insert into tracks(client_id,name) select id,{t(cfg['track'])} from clients where name={t(cfg['client'])} on conflict do nothing;",
       f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,day_column,topic_column)
select c.id,t.id,{t(cfg['toc_file_name'])},{t(cfg['toc_sheet'])},{HR},{q([{'letter': l, 'header': h} for l, h in cols])}::jsonb,'{DC}','{TC}'
from clients c join tracks t on t.client_id=c.id where c.name={t(cfg['client'])} and t.name={t(cfg['track'])}
on conflict (client_id,file_name,sheet_name) do update set columns=excluded.columns, track_id=excluded.track_id;""",
       f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id, x.row_number, x.day_label, x.topic, x.data
from ({ctx}) ctx join toc_files f on f.id=ctx.fid, jsonb_to_recordset({q(toc)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
on conflict (toc_file_id,row_number) do update set day_label=excluded.day_label, topic=excluded.topic, data=excluded.data;"""]
for typ, seq, name, dt, r, rows in items:
    seqv = 'null' if seq is None else t(seq)
    src = q({'source': {'file': 'Details_SkillAssist.xlsx', 'sheet': 'Sony Priyanka', 'row': r}})
    out.append(f"""with ctx as ({ctx}),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'{typ}',{seqv},{t(name)},'{dt}',{src}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,tr.id from ins cross join ctx join toc_rows tr on tr.toc_file_id=ctx.fid and tr.row_number = any(array{rows or [0]}::int[]) on conflict do nothing;""")
out.append('commit;')
print('\n'.join(out))
