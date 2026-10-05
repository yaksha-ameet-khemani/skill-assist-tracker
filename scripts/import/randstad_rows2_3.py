# Randstad sheet of Details_SkillAssist.xlsx, rows 2-3 (user-confirmed 2026-10-05): two ReactJs assignments (20-02-2026)
# above the main table -> track "Module 1: ReactJs" (CSM Aarti) as Daily Assignment Day 4 and Day 5, each linked to its own
# row in the Randstad topic sheet (Topic = ReactJs, Module 1, Subtopics = column C), as randstad.py does for rows 10-25.
# Rows 39-68 ("PL3 HandsOn and Capstones" status table: Level-3 badge tracks, no content names) are skipped (user).
import openpyxl, json, re, warnings
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
FILE, SHEET, CLIENT, TRACK = 'Details_SkillAssist.xlsx', 'Randstad', 'Randstad', 'Module 1: ReactJs'
ws = openpyxl.load_workbook(base + FILE, data_only=True)[SHEET]
one = lambda v: re.sub(r'\s+', ' ', str(v)).strip()
assert one(ws['B2'].value) == 'ReactJs' and one(ws['F2'].value) == '20-02-2026'
q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
out = ['begin;']
for r, day in ((2, '4'), (3, '5')):
    sub, name = one(ws[f'C{r}'].value), one(ws[f'E{r}'].value)
    out.append(f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,{r},null,'ReactJs',{q({'Topic': 'ReactJs', 'Module': 'Module 1', 'Subtopics': sub})}::jsonb from toc_files f join clients c on c.id=f.client_id and c.name={t(CLIENT)}
where f.file_name={t(FILE)} and f.sheet_name={t(SHEET)} on conflict (toc_file_id,row_number) do update set topic=excluded.topic, data=excluded.data;""")
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(TRACK)} join toc_files f on f.client_id=c.id and f.file_name={t(FILE)} and f.sheet_name={t(SHEET)} where c.name={t(CLIENT)}),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'Daily Assignment','{day}',{t(name)},'2026-02-20',{q({'source': {'file': FILE, 'sheet': SHEET, 'row': r}})}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number={r} on conflict do nothing;""")
out.append('commit;')
print('\n'.join(out))
