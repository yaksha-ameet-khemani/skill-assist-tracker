# IBM sheet of Details_SkillAssist.xlsx, rows 23-114 (user-confirmed 2026-10-04): CSM Ayman, track = Phase 1/2/3 only
# (user, 2026-10-04); the sheet's LP Name is kept in extra.lp and its Course column in extra.course (home-page Course column).
# "Milestone Assessment N" -> Milestone Assessment N, "Capstone Project" -> Capstone, everything else -> Daily Assignment
# (incl. rows the sheet marks "Assessment"). No day structure, so sequence is blank except for milestones.
# TOC: Phases 1-2 = IBM/TOC- IBM PJP ... (course level); Phase 3 = IBM/Phase-3_TOC.xlsx, collapsed to one topic per
# course (modules listed in data). Each item links to its course; a milestone to the courses since the previous
# milestone; a capstone to every course of its LP. Rows 2-21 were loaded earlier by hand.
import openpyxl, json, re, warnings, sys, datetime
from collections import OrderedDict
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
FILE, SHEET, CLIENT, CSM = 'Details_SkillAssist.xlsx', 'IBM', 'IBM', 'Ayman'
PJP = 'IBM/TOC- IBM PJP - NEW April - 2026 New onboarding (2).xlsx'
P3 = 'IBM/Phase-3_TOC.xlsx'
MANUAL = '(Added manually)', 'IBM topics'
one = lambda v: re.sub(r'\s+', ' ', str(v).replace('\xa0', ' ')).strip()
key = lambda s: one(s).lower()
# Details course name -> TOC course name where the versions differ
ALIAS = {'Hands-On Object Oriented Programming with Java 11': 'Hands-On Object Oriented Programming with Java 21',
         'Hands-On Enterprise Application Development with Java 9': 'Hands-On Enterprise Application Development with Java 21',
         'The Complete Linux Training Course to Get Your Dream IT Job 2019': 'The Complete Linux Training Course to Get Your Dream IT Job 2025',
         '.Net Framework 4.7': '.Net Framework 4.8', 'Beginning ASP.NET Core 3.0': 'Beginning ASP.NET Core 5.0',
         'AZ-900: Microsoft Azure Fundamentals Certification 2020': 'AZ-900: Microsoft Azure Fundamentals Certification 2025'}
phase = lambda lp: 'Phase 1' if lp.startswith('Phase 1') else 'Phase 2' if lp.startswith('Phase 2') else 'Phase 3'
DATES = {'25-06-2026': '2026-06-25', '21-07-2026': '2026-07-21', '11-08-2026': '2026-08-11',
         '31-08-2026': '2026-08-31', '23-09-2026': '2026-09-23'}

# ---- TOC files: {track: dict(file, sheet, courses=[(row, course, data)], marks=[(pos, 'M'/'C')])}
tocs = OrderedDict()
pjp = openpyxl.load_workbook(base + PJP, data_only=True)
PJP_LP = {'Phase 1:Common Foundation 2023': 'Phase 1:Common Foundation 2023',
          'Phase 2- Developer 2024': 'Phase 2- Developer 2025',
          'Phase 2: Technical Specialist 2024': 'Phase 2: Technical Specialist 2025'}
for sh in ('Phase 1', 'Phase 2'):
    lp = None
    for r in range(3, pjp[sh].max_row + 1):
        a, b, c = (pjp[sh].cell(r, i).value for i in (1, 2, 3))
        if a and one(a) in PJP_LP: lp = PJP_LP[one(a)]
        elif a: lp = None  # block title / header rows
        if not lp or not b or one(b) in ('Course', 'Duration'): continue
        d = tocs.setdefault(lp, dict(file=PJP, sheet=sh, courses=[], marks=[]))
        if one(b).startswith('Milestone Assessment') or one(b) == 'Capstone Project':
            d['marks'].append(len(d['courses']))
        else:
            d['courses'].append((r, one(b), {'LP Name': one(a) if a else lp, 'Course': one(b), 'Course Duration': one(c) if c else None}))

def hms(v):
    if isinstance(v, datetime.time): return v.hour * 3600 + v.minute * 60 + v.second
    m = re.fullmatch(r'(\d+):(\d\d):(\d\d)', str(v or '').strip())
    return int(m[1]) * 3600 + int(m[2]) * 60 + int(m[3]) if m else 0

p3 = openpyxl.load_workbook(base + P3, data_only=True, read_only=True)
P3_SHEETS = {'Application Developer-Cloud': 'Application Developer-Cloud FullStack',
             'Application Developer-Java & We': 'Application Developer-Java & Web Technologies',
             'Quality Engineer-FullStack': 'Quality Engineer-FullStack',
             'Application Developer -FrontEnd': 'Application Developer -Experienced Front End',
             'Application Developer-RDBMS': 'Application Developer-RDBMS',
             'Application Developer-MS .Net': 'Application Developer-Microsoft.Net',
             'Application Developer-DevOps': 'Application Developer-DevOps',
             'Data Engineer-Data Platforms': 'Data Engineer-Data Platforms'}
for sh, track in P3_SHEETS.items():
    lp, courses = None, OrderedDict()
    for r, row in enumerate(p3[sh].iter_rows(min_row=2, max_col=7, values_only=True), start=2):
        a, b, c, d, e, f, g = row
        if a: lp = one(a)
        if lp != track: continue  # the .Net sheet also carries copies of the DevOps and Data Engineer LPs
        if b: cur = courses.setdefault(one(b), dict(row=r, code=one(c) if c else None, modules=OrderedDict(), subs=0, secs=0))
        if d: cur['modules'][one(d)] = 1; cur['subs'] += 1; cur['secs'] += hms(f)
    assert courses, sh
    tocs[track] = dict(file=P3, sheet=sh, marks=[], courses=[
        (x['row'], cn, {'LP Name': track, 'Course': cn, 'Course Code': x['code'], 'Modules': list(x['modules']),
                        'Sub-modules': x['subs'], 'Duration': '%d:%02d:%02d' % (x['secs'] // 3600, x['secs'] // 60 % 60, x['secs'] % 60)})
        for cn, x in courses.items()])

# ---- items from the Details sheet
ws = openpyxl.load_workbook(base + FILE, data_only=True)[SHEET]
items, lp, date = [], None, None
for r in list(range(24, 62)) + list(range(65, 115)):
    a, b, c, d, e = (ws.cell(r, i).value for i in range(1, 6))
    if a: lp = one(a)
    if e: date = DATES[one(e)]
    if not d: continue
    b, c = one(b), one(c)
    if b.startswith('Milestone Assessment'): typ, seq = 'Milestone Assessment', b.split()[-1]
    elif b == 'Capstone Project': typ, seq = 'Capstone', None
    else: typ, seq = 'Daily Assignment', None  # user: everything except milestone/capstone is an assignment
    items.append(dict(row=r, track=lp, type=typ, seq=seq, name=one(d), course=b, date=date))

# ---- links
toc_rows_of = lambda tr: [x[0] for x in tocs[tr]['courses']]
def course_row(tr, course):
    want = key(ALIAS.get(course, course))
    hit = [x[0] for x in tocs[tr]['courses'] if key(x[1]) == want]
    assert len(hit) == 1, (tr, course)
    return hit[0]
for tr in OrderedDict.fromkeys(i['track'] for i in items):
    its = [i for i in items if i['track'] == tr]
    if tr not in tocs:  # Curriculum Open Source: no TOC -> manual topic
        for i in its: i['links'] = ('manual', [8])
        continue
    rows, since = toc_rows_of(tr), []
    marks = tocs[tr]['marks']
    for i in its:
        if i['type'] == 'Capstone': i['links'] = ('toc', rows)
        elif i['type'] == 'Milestone Assessment':
            if marks:  # TOC has its own milestone markers (Phases 1-2)
                n = int(i['seq']); lo = marks[n - 2] if n > 1 else 0
                i['links'] = ('toc', rows[lo:marks[n - 1]])
            else: i['links'] = ('toc', list(OrderedDict.fromkeys(since)))
            since = []
        else:
            cr = course_row(tr, i['course']); i['links'] = ('toc', [cr]); since.append(cr)

# Same content listed twice in one LP (Details rows 83 and 90) -> one item that notes both rows
seen = {}
for i in items:
    k = (i['track'], i['type'], i['seq'], i['name'])
    if k in seen:
        seen[k].setdefault('also_rows', []).append(i['row'])
        seen[k]['links'] = ('toc', list(OrderedDict.fromkeys(seen[k]['links'][1] + i['links'][1]))); i['dup'] = True
    else: seen[k] = i
items = [i for i in items if not i.get('dup')]
for i in items: print(i['row'], i['track'][:28], i['type'], i['seq'] or '', i['name'][:50], len(i['links'][1]), file=sys.stderr)

# ---- SQL
q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
out = ['begin;']
for tr in OrderedDict.fromkeys(phase(i['track']) for i in items):
    out.append(f"insert into tracks(client_id,name,csm) select id,{t(tr)},{t(CSM)} from clients where name={t(CLIENT)} on conflict (client_id,name) do update set csm=excluded.csm;")
files = OrderedDict()
for tr, d in tocs.items(): files.setdefault((d['file'], d['sheet']), []).append(tr)
for (fn, sh), trs in files.items():
    if fn == PJP: cols, hdr, dc, tc = [{'letter': 'A', 'header': 'LP Name'}, {'letter': 'B', 'header': 'Course'}, {'letter': 'C', 'header': 'Course Duration'}], 2, None, 'B'
    else: cols, hdr, dc, tc = [{'letter': 'A', 'header': 'LP Name'}, {'letter': 'B', 'header': 'Course'}, {'letter': 'C', 'header': 'Course Code'}, {'letter': 'D', 'header': 'Modules'}, {'letter': 'E', 'header': 'Sub-modules'}, {'letter': 'F', 'header': 'Duration'}], 1, None, 'B'
    note = None
    if fn == P3: note = 'One topic per course (sub-module rows collapsed; module names kept in data).'
    if sh == 'Application Developer-MS .Net': note += ' This sheet also repeats the DevOps and Data Engineer LPs; only the Application Developer-Microsoft.Net block (row 2782 on) is loaded.'
    if sh == 'Phase 2': note = 'Two LPs in one sheet (Developer, Technical Specialist; TOC says 2024, Details sheet says 2025). Milestone/Capstone rows are markers, not topics.'
    if sh == 'Phase 1': note = 'Milestone/Capstone rows are markers, not topics.'
    trk = f"(select id from tracks where client_id=c.id and name={t(phase(trs[0]))})"
    out.append(f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,day_column,topic_column,notes)
select c.id,{trk},{t(fn.split('/')[-1])},{t(sh)},{hdr},{q(cols)}::jsonb,{t(dc)},{t(tc)},{t(note)} from clients c where c.name={t(CLIENT)}
on conflict (client_id,file_name,sheet_name) do update set notes=excluded.notes, track_id=excluded.track_id;""")
    rows = [dict(row_number=r, day_label=None, topic=cn, data=dd) for tr in trs for r, cn, dd in tocs[tr]['courses']]
    out.append(f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,x.day_label,x.topic,x.data from toc_files f join clients c on c.id=f.client_id and c.name={t(CLIENT)}, jsonb_to_recordset({q(rows)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
where f.file_name={t(fn.split('/')[-1])} and f.sheet_name={t(sh)}
on conflict (toc_file_id,row_number) do update set topic=excluded.topic, data=excluded.data;""")
out.append(f"""insert into toc_rows(toc_file_id,row_number,topic,data)
select f.id,8,'Python with SQL – database CRUD app',{q({'Module': 'Curriculum Open Source', 'Topic': 'Python with SQL – database CRUD app'})}::jsonb
from toc_files f join clients c on c.id=f.client_id and c.name={t(CLIENT)} where f.file_name={t(MANUAL[0])} and f.sheet_name={t(MANUAL[1])}
on conflict (toc_file_id,row_number) do update set topic=excluded.topic, data=excluded.data;""")
for i in items:
    src = {'file': FILE, 'sheet': SHEET, 'row': i['row']}
    extra = {'source': src, 'course': i['course'], 'lp': i['track']}
    if i.get('also_rows'): extra['also_rows'] = i['also_rows']
    kind, rows = i['links']
    fn, sh = (MANUAL if kind == 'manual' else (tocs[i['track']]['file'].split('/')[-1], tocs[i['track']]['sheet']))
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(phase(i['track']))} join toc_files f on f.client_id=c.id and f.file_name={t(fn)} and f.sheet_name={t(sh)} where c.name={t(CLIENT)}),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,{t(i['type'])},{t(i['seq'])},{t(i['name'])},'{i['date']}',{q(extra)}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number = any(array{rows}::int[]) on conflict do nothing;""")
out.append('commit;')
print('\n'.join(out))
