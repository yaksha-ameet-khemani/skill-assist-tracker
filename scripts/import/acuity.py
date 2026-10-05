# Acuity sheet of Details_SkillAssist.xlsx (user-confirmed 2026-10-03): CSM Ayman, 7 Milestone Assessments (M1-M4)
# in 4 tracks, 31-03-2026. TOC = Acuity/17_Day_Training_Plan 1.xlsx: topic = Topics, subtopics = Subtopics,
# plus Day and the TOC's Track. Each milestone links to its track's days; M3 (SQL) only to Data Engineering days 9-11.
import openpyxl, json, re, warnings, sys
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
TFILE, TSHEET = '17_Day_Training_Plan 1.xlsx', '17-Day Training Plan'
FILE, SHEET, CLIENT, CSM, DATE = 'Details_SkillAssist.xlsx', 'Acuity', 'Acuity', 'Ayman', '2026-03-31'
LINKS = {'1': range(3, 6), '2': range(6, 9), '3': range(9, 12), '4': range(15, 17)}   # milestone -> TOC days
one = lambda v: None if v is None or not str(v).strip() else re.sub(r'\s+', ' ', str(v)).strip()

tw = openpyxl.load_workbook(base + 'Acuity/' + TFILE, data_only=True)[TSHEET]
toc, row_of_day, ttrack = [], {}, None
for r in range(2, tw.max_row + 1):
    if one(tw[f'B{r}'].value): ttrack = one(tw[f'B{r}'].value)
    day, topic = one(tw[f'A{r}'].value), one(tw[f'C{r}'].value)
    if not day or not topic: continue
    toc.append(dict(row_number=r, day_label=day, topic=topic,
                    data={'Day': day, 'Track': ttrack, 'Topics': topic, 'Subtopics': one(tw[f'D{r}'].value)}))
    row_of_day[int(day)] = r

ws = openpyxl.load_workbook(base + FILE, data_only=True)[SHEET]
assert ws['D2'].value == '31-03-2026'
items, track, ms = [], None, None
for r in range(2, ws.max_row + 1):
    if one(ws[f'A{r}'].value): track = one(ws[f'A{r}'].value)
    if one(ws[f'B{r}'].value): ms = re.search(r'\d+', ws[f'B{r}'].value).group()
    name = one(ws[f'C{r}'].value)
    if name: items.append(dict(row=r, track=track, seq=ms, name=name, rows=[row_of_day[d] for d in LINKS[ms]]))
for i in items: print(i['row'], i['track'], 'M' + i['seq'], i['name'], '->', [tw[f'C{x}'].value for x in i['rows']], file=sys.stderr)

q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
out = ['begin;', f"insert into clients(name) values ({t(CLIENT)}) on conflict do nothing;"]
for tr in dict.fromkeys(i['track'] for i in items):
    out.append(f"insert into tracks(client_id,name,csm) select id,{t(tr)},{t(CSM)} from clients where name={t(CLIENT)} on conflict (client_id,name) do update set csm=excluded.csm;")
cols = [{'letter': 'A', 'header': 'Day'}, {'letter': 'B', 'header': 'Track'}, {'letter': 'C', 'header': 'Topics'}, {'letter': 'D', 'header': 'Subtopics'}]
out.append(f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns,day_column,topic_column)
select id,null,{t(TFILE)},{t(TSHEET)},1,{q(cols)}::jsonb,'A','C' from clients where name={t(CLIENT)}
on conflict (client_id,file_name,sheet_name) do update set columns=excluded.columns;""")
out.append(f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,x.day_label,x.topic,x.data from toc_files f join clients c on c.id=f.client_id and c.name={t(CLIENT)}, jsonb_to_recordset({q(toc)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
where f.file_name={t(TFILE)} and f.sheet_name={t(TSHEET)}
on conflict (toc_file_id,row_number) do update set day_label=excluded.day_label, topic=excluded.topic, data=excluded.data;""")
for i in items:
    extra = q({'source': {'file': FILE, 'sheet': SHEET, 'row': i['row']}})
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(i['track'])} join toc_files f on f.client_id=c.id and f.file_name={t(TFILE)} and f.sheet_name={t(TSHEET)} where c.name={t(CLIENT)}),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'Milestone Assessment',{t(i['seq'])},{t(i['name'])},'{DATE}',{extra}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number = any(array{i['rows']}::int[]) on conflict do nothing;""")
out.append('commit;')
print('\n'.join(out))
