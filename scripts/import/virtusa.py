# Virtusa sheet of Details_SkillAssist.xlsx (user-confirmed 2026-10-03): track SQL, CSM Priyanka, all Daily
# Assignments, 30-04-2026 (merged). No Day column, so days are numbered in row order. No TOC: one manual topic each.
import openpyxl, json, re, warnings, sys
warnings.filterwarnings('ignore')
base = '/home/administrator/Downloads/IntelliJ Projects/Skill Assist Tracker/'
FILE, SHEET, CLIENT, TRACK, CSM, DATE = 'Details_SkillAssist.xlsx', 'Virtusa', 'Virtusa', 'SQL', 'Priyanka', '2026-04-30'
MANUAL = '(Added manually)', 'Virtusa topics'
TOPICS = {2: 'SQL Date Functions & Filtering', 3: 'SQL Aggregation & GROUP BY', 4: 'SQL Joins & Aggregation'}
ws = openpyxl.load_workbook(base + FILE, data_only=True)[SHEET]
one = lambda v: re.sub(r'\s+', ' ', str(v)).strip()
assert ws['B2'].value == '30-04-2026'
toc = [dict(row_number=r, day_label=None, topic=tp, data={'Topic': tp}) for r, tp in TOPICS.items()]
items = [dict(row=r, seq=str(n), name=one(ws[f'A{r}'].value), rows=[r]) for n, r in enumerate(TOPICS, 1)]
for i in items: print(i['seq'], i['name'], '->', TOPICS[i['row']], file=sys.stderr)

q = lambda o: '$j$' + json.dumps(o, ensure_ascii=False) + '$j$'
t = lambda s: 'null' if s is None else '$n$' + s + '$n$'
out = ['begin;', f"insert into clients(name) values ({t(CLIENT)}) on conflict do nothing;",
       f"insert into tracks(client_id,name,csm) select id,{t(TRACK)},{t(CSM)} from clients where name={t(CLIENT)} on conflict (client_id,name) do update set csm=excluded.csm;",
       f"""insert into toc_files(client_id,track_id,file_name,sheet_name,header_row,columns)
select id,null,{t(MANUAL[0])},{t(MANUAL[1])},1,{q([{'letter': 'A', 'header': 'Topic'}])}::jsonb from clients where name={t(CLIENT)}
on conflict (client_id,file_name,sheet_name) do nothing;""",
       f"""insert into toc_rows(toc_file_id,row_number,day_label,topic,data)
select f.id,x.row_number,x.day_label,x.topic,x.data from toc_files f join clients c on c.id=f.client_id and c.name={t(CLIENT)}, jsonb_to_recordset({q(toc)}::jsonb) as x(row_number int, day_label text, topic text, data jsonb)
where f.file_name={t(MANUAL[0])} and f.sheet_name={t(MANUAL[1])}
on conflict (toc_file_id,row_number) do update set topic=excluded.topic, data=excluded.data;"""]
for i in items:
    extra = q({'source': {'file': FILE, 'sheet': SHEET, 'row': i['row']}})
    out.append(f"""with ctx as (select c.id cid, tr.id tid, f.id fid from clients c join tracks tr on tr.client_id=c.id and tr.name={t(TRACK)} join toc_files f on f.client_id=c.id and f.file_name={t(MANUAL[0])} and f.sheet_name={t(MANUAL[1])} where c.name={t(CLIENT)}),
ins as (insert into contents(client_id,track_id,content_type,sequence_label,name,delivery_date,extra)
 select cid,tid,'Daily Assignment',{t(i['seq'])},{t(i['name'])},'{DATE}',{extra}::jsonb from ctx
 on conflict (client_id,track_id,content_type,sequence_label,name) do update set delivery_date=excluded.delivery_date, extra=excluded.extra returning id)
insert into content_toc_links(content_id,toc_row_id)
select ins.id,r.id from ins cross join ctx join toc_rows r on r.toc_file_id=ctx.fid and r.row_number = any(array{i['rows']}::int[]) on conflict do nothing;""")
out.append('commit;')
print('\n'.join(out))
